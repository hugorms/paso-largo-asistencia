import asyncio
import base64
import io
import logging
import re
import uuid
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from fastapi import FastAPI, Request
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook
from starlette.middleware.sessions import SessionMiddleware

from . import alcance, attendance, db, devices, reports, seguridad
from .isapi import MINOR_FAIL, MINOR_NAMES, ISAPIError

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
logging.getLogger("httpx").setLevel(logging.WARNING)
SECRET = db.init()

app = FastAPI(title="Paso Largo", docs_url=None, redoc_url=None, openapi_url=None)
app.mount("/static", StaticFiles(directory=str(Path(__file__).parent / "static")), name="static")
app.add_middleware(SessionMiddleware, secret_key=SECRET, max_age=12 * 3600, same_site="lax")
# Va después de la sesión para quedar por fuera: CSRF, cabeceras de seguridad y cookie segura por HTTPS
app.middleware("http")(seguridad.middleware)
templates = Jinja2Templates(directory=str(Path(__file__).parent / "templates"))
templates.env.globals.update(local_offset=devices.local_offset, hm=attendance.hm, colors=attendance.STATUS_COLORS, minor_names=MINOR_NAMES,
                             days=reports.DAYS, today=date.today, clock_warnings=devices.CLOCK_WARNINGS,
                             server_now=lambda: datetime.now().strftime("%d/%m/%Y %H:%M"))

# En el servidor la imagen trae /app/VERSION con la fecha del despliegue; en la PC no existe
_version_file = Path(__file__).parent.parent / "VERSION"
templates.env.globals["version"] = _version_file.read_text().strip() if _version_file.exists() else "local"

DEFAULT_PASSWORD = "admin123"
MIN_PASSWORD = 8
ROLES = alcance.ROLES
LEAVE_KINDS = ["Vacaciones", "Permiso médico", "Permiso personal", "Licencia", "Comisión de servicio", "Otro"]
poller = devices.Poller()


@app.on_event("startup")
def _start():
    poller.start()


# ------------------------------------------------------------------ auth helpers
class Redirect(Exception):
    def __init__(self, url):
        self.url = url


@app.exception_handler(Redirect)
async def _redir(request, exc):
    return RedirectResponse(exc.url, status_code=303)


def user(request: Request, *roles):
    uid = request.session.get("uid")
    u = db.q1("SELECT * FROM users WHERE id=? AND active=1", (uid,)) if uid else None
    if not u or request.session.get("pv") != seguridad.password_tag(SECRET, u["password_hash"]):
        # sin sesión, o la contraseña cambió desde que entró: hay que volver a ingresar
        request.session.pop("uid", None)
        raise Redirect("/login")
    if request.session.get("cambiar_clave") and request.url.path not in ("/cuenta", "/logout"):
        flash(request, "Antes de seguir, cambia la contraseña por defecto.", "warning")
        raise Redirect("/cuenta")
    if (not alcance.is_admin(u) and request.url.path not in ("/cuenta", "/logout", "/sin-local")
            and not alcance.allowed(u)):
        # encargado o consulta sin ningún local activo asignado: no hay nada que mostrarle todavía
        raise Redirect("/sin-local")
    if roles and u["role"] not in roles:
        flash(request, "No tienes permiso para esa sección.", "danger")
        raise Redirect("/")
    return u


def flash(request, msg, kind="success"):
    # se reasigna (no .append) para que la sesión siempre se dé por modificada y el aviso no se pierda
    request.session["flash"] = request.session.get("flash", []) + [[kind, msg]]


def render(request, name, u, **ctx):
    msgs = request.session.pop("flash", [])
    return templates.TemplateResponse(request, name, {
        "u": u, "flashes": msgs, "roles": ROLES, "company": db.get_setting("company", ""),
        "path": request.url.path, "local_choices": alcance.choices(u), "local_now": alcance.current(request, u),
        **ctx})


def back(url):
    return RedirectResponse(url, status_code=303)


async def form(request):
    return {k: (v.strip() if isinstance(v, str) else v) for k, v in (await request.form()).multi_items()}


# ------------------------------------------------------------------ login
@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return templates.TemplateResponse(request, "login.html", {"error": None, "company": db.get_setting("company"),
                                                              "flashes": request.session.pop("flash", [])})


_DUMMY_HASH = db.hash_password(uuid.uuid4().hex)


@app.post("/login")
async def login(request: Request):
    f = await form(request)
    username, ip = f.get("username", "")[:64], seguridad.client_ip(request)
    error = lambda msg, status=200: templates.TemplateResponse(
        request, "login.html", {"error": msg, "company": db.get_setting("company")}, status_code=status)
    wait = seguridad.login_blocked(ip, username)
    if wait:
        return error(f"Demasiados intentos fallidos. Espera {wait} minutos y vuelve a intentar.", 429)
    u = db.q1("SELECT * FROM users WHERE username=? AND active=1", (username,))
    # Se compara igual aunque el usuario no exista, para no delatar qué usuarios existen
    good = await run_in_threadpool(db.check_password, f.get("password", ""), u["password_hash"] if u else _DUMMY_HASH)
    if not u or not good:
        seguridad.login_failed(ip, username)
        db.audit(username, "ingreso fallido", ip)
        await asyncio.sleep(1)
        return error("Usuario o contraseña incorrectos")
    seguridad.login_ok(ip, username)
    request.session.clear()
    request.session["uid"] = u["id"]
    request.session["pv"] = seguridad.password_tag(SECRET, u["password_hash"])
    db.audit(u["username"], "login", ip)
    if f.get("password") == DEFAULT_PASSWORD:
        request.session["cambiar_clave"] = True
        return back("/cuenta")
    return back("/resumen" if alcance.is_admin(u) else "/")


@app.get("/sin-local", response_class=HTMLResponse)
def no_local(request: Request):
    u = user(request)
    if alcance.is_admin(u) or alcance.allowed(u):
        return back("/")
    return templates.TemplateResponse(request, "sin_local.html", {"u": u, "roles": ROLES})


@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return back("/login")


@app.post("/local")
async def local_choose(request: Request):
    """Selector del menú: qué local se está viendo (vacío = todos los permitidos)."""
    u, f = user(request), await form(request)
    alcance.choose(request, u, f.get("local"))
    # next: a dónde ir después (p. ej. «Ver detalle» en el resumen de locales); solo rutas de este sistema
    fallback = seguridad.safe_back(request.headers.get("referer"), "/")
    nxt = (f.get("next") or "").replace("\\", "/")  # «/\otro.sitio» el navegador lo toma como «//otro.sitio»
    return back(seguridad.safe_back(nxt, fallback) if nxt else fallback)


@app.get("/cuenta", response_class=HTMLResponse)
def account(request: Request):
    return render(request, "account.html", user(request))


@app.post("/cuenta")
async def account_save(request: Request):
    u, f = user(request), await form(request)
    if not db.check_password(f.get("current", ""), u["password_hash"]):
        flash(request, "La contraseña actual no es correcta.", "danger")
    elif len(f.get("new", "")) < MIN_PASSWORD or f["new"] != f.get("confirm"):
        flash(request, f"La nueva contraseña debe tener {MIN_PASSWORD}+ caracteres y coincidir.", "danger")
    elif f["new"] == DEFAULT_PASSWORD or f["new"] == f.get("current"):
        flash(request, "Elige una contraseña distinta a la anterior.", "danger")
    else:
        new_hash = db.hash_password(f["new"])
        db.ex("UPDATE users SET password_hash=? WHERE id=?", (new_hash, u["id"]))
        # Esta sesión sigue abierta; las demás sesiones de este usuario se cierran
        request.session["pv"] = seguridad.password_tag(SECRET, new_hash)
        request.session.pop("cambiar_clave", None)
        db.audit(u["username"], "cambio de contraseña")
        flash(request, "Contraseña actualizada.")
    return back("/cuenta")


# ------------------------------------------------------------------ dashboard
TREND_KEYS = {"Puntual": "puntual", "Tarde": "tarde", "Ausente": "ausente"}


def emp_locales(request, u):
    """Locales de los empleados que se ven ahora, para attendance.compute. El admin sin elegir local ve los
    locales activos y también a los empleados que aún no tienen local (None en la lista)."""
    vis = alcance.ids(request, u)
    if alcance.is_admin(u) and request.session.get("local") not in vis:
        return vis + [None]
    return vis


def trend_data(days: int, locales=None):
    """Puntuales, tardanzas y faltas de los últimos `days` días (0 = últimas 24 h, por hora),
    más los totales del periodo anterior para comparar. locales: como en attendance.compute."""
    today = date.today()
    if not days:
        now = datetime.now().replace(minute=0, second=0, microsecond=0)
        first = now - timedelta(hours=23)
        by_hour = defaultdict(lambda: {"puntual": 0, "tarde": 0, "ausente": 0})
        rows = attendance.compute(today - timedelta(days=2), today, locales=locales)
        for r in rows:
            if r["entry"] and r["status"] in ("Puntual", "Tarde"):
                by_hour[r["entry"].replace(minute=0, second=0, microsecond=0)][TREND_KEYS[r["status"]]] += 1
        series = [{"date": h.strftime("%H:00"), "dow": "", **by_hour[h]} for h in (first + timedelta(hours=i) for i in range(24))]
        total = {k: sum(x[k] for x in series) for k in ("puntual", "tarde")}
        prev = {k: sum(by_hour[first - timedelta(hours=i)][k] for i in range(1, 25)) for k in ("puntual", "tarde")}
        # las faltas no tienen hora: se comparan las de hoy con las de ayer
        total["ausente"] = sum(r["status"] == "Ausente" for r in rows if r["date"] == today)
        prev["ausente"] = sum(r["status"] == "Ausente" for r in rows if r["date"] == today - timedelta(days=1))
        return {"days": 0, "series": series, "total": total, "prev": prev}
    start = today - timedelta(days=days - 1)
    by_day = defaultdict(lambda: {"puntual": 0, "tarde": 0, "ausente": 0})
    for r in attendance.compute(start - timedelta(days=days), today, locales=locales):
        if r["status"] in TREND_KEYS:
            by_day[r["date"]][TREND_KEYS[r["status"]]] += 1
    series = [{"date": d.strftime("%d/%m"), "dow": "LMMJVSD"[d.weekday()], **by_day[d]}
              for d in (start + timedelta(days=i) for i in range(days))]
    total = {k: sum(x[k] for x in series) for k in ("puntual", "tarde", "ausente")}
    prev = {k: sum(by_day[start - timedelta(days=i)][k] for i in range(1, days + 1)) for k in ("puntual", "tarde", "ausente")}
    return {"days": days, "series": series, "total": total, "prev": prev}


def _mark_away(request, u, evs):
    """Prepara eventos para mostrar: marca los hechos en un equipo de otro local que el del empleado y
    esconde quién es cuando el empleado es de un local que este usuario no puede ver."""
    for ev in evs:
        if ev.get("emp_id") and not alcance.can_see(request, u, ev.get("emp_local")):
            ev.update(emp_name="Persona de otro local", name=None, employee_no=None, emp_id=None, emp_local=None,
                      picture_url=None, department=None, home=None, other=True)
        ev["away"] = ev.get("dev_local") if attendance.away_from(ev.get("home"), ev.get("dev_local")) else None
    return evs


def visible_events(request, u, where, args, limit, extra_select="", extra_join=""):
    """Eventos de los empleados que se ven ahora (en cualquier equipo) y de los equipos que se ven ahora
    (de cualquier persona; los de otro local salen sin nombre ni foto)."""
    ce, ae = alcance.sql(request, u, "e.local_id", include_null=True)
    cd, ad = alcance.sql(request, u, "d.local_id", include_null=True)
    rows = db.q(f"""SELECT ev.*, e.name AS emp_name, e.id AS emp_id, e.local_id AS emp_local, le.name AS home,
                    d.name AS device, ld.name AS dev_local, dp.name AS department{extra_select} FROM events ev
                    LEFT JOIN employees e ON e.employee_no=ev.employee_no LEFT JOIN devices d ON d.id=ev.device_id
                    LEFT JOIN departments dp ON dp.id=e.department_id LEFT JOIN locales le ON le.id=e.local_id
                    LEFT JOIN locales ld ON ld.id=d.local_id {extra_join}
                    WHERE {' AND '.join(list(where) + [f"({ce} OR {cd})"])}
                    ORDER BY ev.event_time DESC, ev.id DESC LIMIT {int(limit)}""", list(args) + ae + ad)
    return _mark_away(request, u, rows)


def dashboard_data(request, u, days: int = 7):
    today = date.today()
    locs = emp_locales(request, u)
    rows = attendance.compute(today, today, locales=locs)
    count = lambda st: sum(1 for r in rows if r["status"] == st)
    fails = ",".join(map(str, sorted(MINOR_FAIL)))
    last = visible_events(request, u, [f"(ev.employee_no IS NOT NULL OR ev.minor IN ({fails}))"], [], 15)
    for ev in last:
        ev["minor_name"] = MINOR_NAMES.get(ev["minor"], f"Evento {ev['minor']}") if ev["source"] == "equipo" else "Marca manual"
    cd, ad = alcance.sql(request, u, "local_id", include_null=True)
    devs = db.q(f"SELECT id, name, online, last_seen, last_error FROM devices WHERE active=1 AND {cd}", ad)
    return {"total": db.q1(f"SELECT COUNT(*) n FROM employees WHERE active=1 AND {cd}", ad)["n"],
            "puntual": count("Puntual"), "tarde": count("Tarde"), "ausente": count("Ausente"),
            "pendiente": count("Pendiente"), "justificado": count("Justificado"),
            "presentes": sum(1 for r in rows if r["entry"]),
            "late_list": [{"name": r["name"], "entry": r["entry"].strftime("%H:%M"), "late": r["late_min"]}
                          for r in rows if r["status"] == "Tarde"],
            "absent_list": [r["name"] for r in rows if r["status"] in ("Ausente", "Pendiente")],
            "away_list": [{"name": r["name"], "department": r["department"], "home": r["home_sede"],
                           "sedes": ", ".join(r["away"]), "back": not attendance.away_from(r["home_sede"], r["trail"][-1][0]),
                           "trail": [{"sede": s, "time": t.strftime("%H:%M")} for s, t in r["trail"]]}
                          for r in rows if r["away"]],
            "trend": trend_data(days, locs), "events": last, "devices": devs, "now": datetime.now().strftime("%H:%M:%S")}


@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request):
    u = user(request)
    return render(request, "dashboard.html", u, d=dashboard_data(request, u))


@app.get("/api/dashboard")
def api_dashboard(request: Request, dias: int = 7):
    u = user(request)
    return JSONResponse(dashboard_data(request, u, dias if dias in (0, 7, 30) else 7))


# ------------------------------------------------------------------ empleados
def emp_lookups(local_ids):
    """Áreas y horarios de esos locales (cada uno lleva su local_id para filtrarlos en el formulario)."""
    if not local_ids:
        return {"departments": [], "schedules": []}
    marks = ",".join("?" * len(local_ids))
    return {"departments": db.q(f"""SELECT d.*, l.name AS local FROM departments d JOIN locales l ON l.id=d.local_id
                                    WHERE d.local_id IN ({marks}) ORDER BY l.name, d.name""", local_ids),
            "schedules": db.q(f"SELECT * FROM schedules WHERE local_id IN ({marks}) ORDER BY name", local_ids)}


def next_employee_no():
    """Siguiente N° libre en toda la empresa (el N° es único entre todos los locales)."""
    return str((db.q1("SELECT MAX(CAST(employee_no AS INTEGER)) m FROM employees")["m"] or 0) + 1)


def _emp_form(request, u, e, sync=()):
    """Formulario de empleado. Si no se sabe a qué local va (o se puede cambiar), muestra el selector
    con solo los locales permitidos; áreas y horarios solo de esos locales."""
    allowed = alcance.allowed(u)
    if e.get("id"):
        # al editar se puede cambiar a otro local permitido, nunca a uno ajeno
        locs = allowed if (e.get("local_id") in allowed or alcance.is_admin(u)) else []
    else:
        locs = [i for i in alcance.ids(request, u) if i in allowed]
    pick = len(locs) > 1 or (bool(locs) and e.get("local_id") not in locs)
    lookup_ids = locs or ([e["local_id"]] if e.get("local_id") is not None else [])
    local_opts = [c for c in alcance.choices(u) if c["id"] in locs]
    return render(request, "employee_form.html", u, e=e, sync=sync, pick_local=pick, local_opts=local_opts,
                  **emp_lookups(lookup_ids))


@app.get("/empleados", response_class=HTMLResponse)
def employees(request: Request, q: str = "", dep: str = "", inactivos: str = ""):
    u = user(request)
    cond, args = alcance.sql(request, u, "e.local_id", include_null=True)
    where = [cond]
    if not inactivos:
        where.append("e.active=1")
    if q:
        where.append("(e.name LIKE ? OR e.employee_no LIKE ? OR e.document LIKE ?)")
        args += [f"%{q}%"] * 3
    if dep:
        where.append("e.department_id=?")
        args.append(dep)
    emps = db.q(f"""SELECT e.*, d.name AS department, s.name AS schedule,
                    (SELECT COUNT(*) FROM employee_sync x WHERE x.employee_id=e.id AND x.status='ok') AS sync_ok,
                    (SELECT COUNT(*) FROM employee_sync x WHERE x.employee_id=e.id AND x.status='error') AS sync_err
                    FROM employees e LEFT JOIN departments d ON d.id=e.department_id
                    LEFT JOIN schedules s ON s.id=e.schedule_id WHERE {' AND '.join(where)} ORDER BY e.name""", args)
    lids = alcance.ids(request, u)
    dcond, dargs = alcance.sql(request, u, "local_id")
    return render(request, "employees.html", u, emps=emps, q=q, dep=dep, inactivos=inactivos,
                  devs=db.q(f"SELECT id, name, host FROM devices WHERE active=1 AND {dcond} ORDER BY name", dargs),
                  local_names={r["id"]: r["name"] for r in db.q("SELECT id, name FROM locales")} if len(lids) > 1 else None,
                  local_opts=alcance.choices(u) if alcance.default_local(request, u) is None else None,
                  **emp_lookups(lids))


@app.get("/empleados/nuevo", response_class=HTMLResponse)
def employee_new(request: Request):
    u = user(request, "admin", "encargado")
    return _emp_form(request, u, {"employee_no": next_employee_no(), "active": 1,
                                  "local_id": alcance.default_local(request, u)})


def _emp_visible(request, u, eid):
    """El empleado si este usuario puede verlo; si no (otro local o no existe), vuelve al listado con aviso."""
    e = alcance.row(request, u, "employees", eid)
    if not e:
        flash(request, "Ese empleado no existe o es de otro local.", "danger")
        raise Redirect("/empleados")
    return e


@app.get("/empleados/{eid}", response_class=HTMLResponse)
def employee_edit(request: Request, eid: int):
    u = user(request, "admin", "encargado")
    e = _emp_visible(request, u, eid)
    sync = [s for s in db.q("""SELECT s.*, d.name AS device, d.local_id FROM employee_sync s
                               JOIN devices d ON d.id=s.device_id WHERE s.employee_id=?""", (eid,))
            if alcance.can_see(request, u, s["local_id"])]  # equipos de otros locales: ni el nombre
    return _emp_form(request, u, e, sync)


@app.post("/empleados/guardar")
async def employee_save(request: Request):
    u = user(request, "admin", "encargado")
    raw = await request.form(max_part_size=8 * 1024 * 1024)  # photo_data (foto en texto) puede pasar de 1 MB
    f = {k: (v.strip() if isinstance(v, str) else v) for k, v in raw.items()}
    eid = int(f["id"]) if (f.get("id") or "").isdigit() else None
    old = _emp_visible(request, u, eid) if eid else None
    no = f.get("employee_no", "")

    def fail(msg):
        # vuelve a mostrar el formulario con lo que escribió (la foto hay que volver a elegirla)
        flash(request, msg, "danger")
        e = {k: f.get(k) for k in ("employee_no", "name", "document", "position", "email", "phone", "gender",
                                   "hire_date")}
        e.update(id=eid, active=1 if f.get("active") else 0, photo=old and old["photo"], face_at=old and old["face_at"],
                 department_id=int(f["department_id"]) if (f.get("department_id") or "").isdigit() else None,
                 schedule_id=int(f["schedule_id"]) if (f.get("schedule_id") or "").isdigit() else None,
                 local_id=lid)
        return _emp_form(request, u, e)

    # Local del empleado: el elegido en el formulario, o el que ya tenía, o el que se está viendo
    lid = int(f["local_id"]) if (f.get("local_id") or "").isdigit() else (
        old["local_id"] if old else alcance.default_local(request, u))
    if lid is None or lid not in alcance.allowed(u):
        lid = old["local_id"] if old else None
        return fail("Elige el local del empleado (solo puedes usar tus locales).")
    if not re.fullmatch(r"[A-Za-z0-9]{1,32}", no) or not f.get("name"):
        return fail("El N° de empleado (solo letras/números) y el nombre son obligatorios.")
    # El N° es único en toda la empresa: si es de otro local no se dice de quién
    dup = db.q1("SELECT id, name, local_id FROM employees WHERE employee_no=? AND id<>?", (no, eid or 0))
    if dup:
        nxt = next_employee_no()
        if alcance.can_see(request, u, dup["local_id"]):
            return fail(f"Ya existe un empleado con el N° {no} ({dup['name']}). El siguiente N° libre es {nxt}.")
        return fail(f"Ese N° ya lo usa otro local. El siguiente N° libre es {nxt}.")
    # Áreas y horarios: solo los del local del empleado
    if f.get("department_id") and not db.q1("SELECT 1 FROM departments WHERE id=? AND local_id=?",
                                            (f["department_id"], lid)):
        return fail("El área elegida no es de ese local.")
    if f.get("schedule_id") and not db.q1("SELECT 1 FROM schedules WHERE id=? AND local_id=?", (f["schedule_id"], lid)):
        return fail("El horario elegido no es de ese local.")

    photo_name = None
    data = raw.get("photo_data")
    upload = next((raw.get(k) for k in ("photo", "photo_cam") if getattr(raw.get(k), "filename", "")), None)
    if (isinstance(data, str) and data) or upload is not None:
        try:
            if isinstance(data, str) and data:
                # Foto ya achicada en el navegador: "data:image/jpeg;base64,..."
                content = base64.b64decode(data.split(",", 1)[-1])
            else:
                content = await upload.read()
            jpeg = devices.prepare_face(content)
        except Exception:
            return fail("La foto no es una imagen válida.")
        photo_name = f"{no}_{uuid.uuid4().hex[:8]}.jpg"
        (db.PHOTOS_DIR / photo_name).write_bytes(jpeg)

    vals = [no, f["name"], f.get("document"), f.get("department_id") or None, f.get("position"), f.get("email"),
            f.get("phone"), f.get("gender") or "unknown", f.get("schedule_id") or None, f.get("hire_date") or None,
            1 if f.get("active") else 0, lid]
    renamed = moved_from = None
    if eid:
        if old["employee_no"] != no:
            renamed = old
            # Sus marcas anteriores pasan al N° nuevo, para que el historial siga junto
            db.ex("UPDATE events SET employee_no=? WHERE employee_no=?", (no, old["employee_no"]))
        if old["local_id"] is not None and old["local_id"] != lid:
            moved_from = old["local_id"]
        db.ex("""UPDATE employees SET employee_no=?, name=?, document=?, department_id=?, position=?, email=?,
                 phone=?, gender=?, schedule_id=?, hire_date=?, active=?, local_id=? WHERE id=?""", vals + [eid])
        if photo_name:
            db.ex("UPDATE employees SET photo=?, face_at=COALESCE(face_at, ?) WHERE id=?",
                  (photo_name, date.today().isoformat(), eid))
            if old["photo"]:
                (db.PHOTOS_DIR / old["photo"]).unlink(missing_ok=True)
    else:
        eid = db.ex("""INSERT INTO employees(employee_no, name, document, department_id, position, email, phone,
                       gender, schedule_id, hire_date, active, local_id, photo, face_at)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    vals + [photo_name, date.today().isoformat() if photo_name else None])
    db.audit(u["username"], "guardar empleado", f"{no} {f['name']}")
    emp = db.q1("SELECT * FROM employees WHERE id=?", (eid,))
    if moved_from and not renamed:
        # Pasó a otro local: se retira de los equipos del local anterior (si cambió el N° ya se borra de todos abajo)
        for name, ok, m in await run_in_threadpool(devices.remove_employee, old, [moved_from]):
            if not ok:
                flash(request, f"{name}: no se pudo retirar del equipo del local anterior ({m}).", "warning")
        db.audit(u["username"], "empleado cambia de local", f"{no}: {moved_from} -> {lid}")
    if renamed:
        # Se borra del equipo el usuario con el N° viejo; si no, el equipo lo seguiría reconociendo
        # y volvería a aparecer aquí como otro empleado
        for name, ok, m in await run_in_threadpool(devices.remove_employee, renamed):
            if not ok:
                flash(request, f"{name}: no se pudo borrar el N° viejo {renamed['employee_no']} del equipo ({m}). "
                               "Bórralo a mano en el equipo.", "warning")
        db.audit(u["username"], "cambio de N° de empleado", f"{renamed['employee_no']} -> {no}")
    if emp["active"]:
        res = await run_in_threadpool(devices.push_employee, emp)
    else:
        res = await run_in_threadpool(devices.remove_employee, emp)
    _flash_results(request, "Empleado guardado.", res)
    return back(f"/empleados/{eid}")


def _flash_results(request, msg, res):
    u = user(request)
    if not alcance.is_admin(u):
        mine = {d["name"] for d in db.q(f"SELECT name FROM devices WHERE {alcance.sql(request, u, 'local_id')[0]}",
                                         alcance.sql(request, u, "local_id")[1])}
        others = [r for r in res if r[0] not in mine]
        res = [r for r in res if r[0] in mine]
        if others:
            flash(request, f"También se actualizó en {len(others)} equipo(s) de otros locales.", "info")
    if not res:
        flash(request, msg + " (Su local no tiene equipos activos.)", "warning")
    for name, ok, m in res:
        flash(request, f"{msg} {name}: {m}", "success" if ok else "danger")
        msg = ""


@app.post("/empleados/{eid}/sincronizar")
async def employee_sync(request: Request, eid: int):
    u = user(request, "admin", "encargado")
    emp = _emp_visible(request, u, eid)
    _flash_results(request, "", await run_in_threadpool(devices.push_employee, emp))
    return back(f"/empleados/{eid}")


@app.post("/empleados/{eid}/eliminar")
async def employee_delete(request: Request, eid: int):
    u = user(request, "admin", "encargado")
    emp = _emp_visible(request, u, eid)
    # Se borra de los equipos de todos los locales: ese N° es solo de esta persona (ver devices.remove_employee)
    res = await run_in_threadpool(devices.remove_employee, emp)
    # Se conserva el historial de marcas; el empleado queda inactivo
    db.ex("UPDATE employees SET active=0 WHERE id=?", (eid,))
    db.audit(u["username"], "desactivar empleado", f"{emp['employee_no']} {emp['name']}")
    _flash_results(request, "Empleado desactivado y retirado de los equipos.", res)
    return back("/empleados")


@app.post("/empleados/sincronizar-todos")
async def employees_sync_all(request: Request):
    u = user(request, "admin", "encargado")
    ok = err = 0
    # Solo los empleados de los locales que se están viendo; cada uno va a los equipos de su local
    cond, args = alcance.sql(request, u, "local_id")
    for emp in db.q(f"SELECT * FROM employees WHERE active=1 AND {cond}", args):
        for _, good, _m in await run_in_threadpool(devices.push_employee, emp):
            ok, err = ok + good, err + (not good)
    db.audit(u["username"], "sincronizar todos", ", ".join(str(i) for i in args))
    flash(request, f"Sincronización terminada: {ok} correctas, {err} con error.", "success" if not err else "warning")
    return back("/empleados")


@app.post("/empleados/importar-equipo")
async def employees_import_device(request: Request):
    u = user(request, "admin", "encargado")
    f = await form(request)
    dev = alcance.row(request, u, "devices", f.get("device_id")) if (f.get("device_id") or "").isdigit() else None
    if not dev or dev["local_id"] is None:
        flash(request, "Elige un equipo de tu local.", "danger")
        return back("/empleados")
    try:
        users = await run_in_threadpool(devices.client(dev).search_users)
    except ISAPIError as e:
        flash(request, str(e), "danger")
        return back("/empleados")
    # Quedan en el local del equipo, con un horario de ese local; los N° que ya existen (en cualquier local) no se tocan
    n = 0
    for us in users:
        no = str(us.get("employeeNo") or "").strip()
        if no and not db.q1("SELECT 1 FROM employees WHERE employee_no=?", (no,)):
            devices.add_from_device(dev, no, us.get("name") or no, us.get("gender", "unknown"))
            n += 1
    db.audit(u["username"], "importar del equipo", f"{dev['name']}: {n}")
    flash(request, f"Se importaron {n} empleados nuevos desde {dev['name']} ({len(users)} en el equipo). "
                   "Las fotos se quedan en el equipo; súbelas aquí si quieres verlas en el sistema.")
    return back("/empleados")


@app.post("/empleados/importar-excel")
async def employees_import_excel(request: Request):
    u = user(request, "admin", "encargado")
    raw = await request.form()
    # Local donde quedan los empleados nuevos y las áreas nuevas: el que se ve, o el elegido en el formulario
    lid = int(raw["local_id"]) if str(raw.get("local_id") or "").isdigit() else alcance.default_local(request, u)
    if lid is None or lid not in alcance.allowed(u):
        flash(request, "Elige el local donde se importan los empleados.", "danger")
        return back("/empleados")
    try:
        wb = load_workbook(io.BytesIO(await raw["file"].read()), read_only=True)
    except Exception:
        flash(request, "No se pudo leer el archivo. Usa un .xlsx.", "danger")
        return back("/empleados")
    rows = list(wb.active.iter_rows(values_only=True))
    if not rows:
        flash(request, "El Excel está vacío.", "danger")
        return back("/empleados")
    head = [str(h or "").strip().lower() for h in rows[0]]
    col = lambda *names: next((head.index(n) for n in names if n in head), None)
    c_no, c_name = col("numero", "número", "n°", "codigo", "código", "employee_no"), col("nombre", "name")
    c_doc, c_dep, c_pos = col("documento", "dni", "cedula", "cédula"), col("departamento", "area", "área"), col("cargo", "puesto")
    c_mail, c_tel = col("email", "correo"), col("telefono", "teléfono", "celular")
    if c_no is None or c_name is None:
        flash(request, "El Excel debe tener columnas 'numero' y 'nombre'.", "danger")
        return back("/empleados")
    sched = devices.local_schedule(lid)
    n, other = 0, []
    for r in rows[1:]:
        g = lambda c: (str(r[c]).strip() if c is not None and r[c] is not None else None)
        no, name = g(c_no), g(c_name)
        if not no or not name:
            continue
        no = no.removesuffix(".0")
        cur = db.q1("SELECT id, local_id FROM employees WHERE employee_no=?", (no,))
        if cur and not alcance.can_see(request, u, cur["local_id"]):
            other.append(no)  # ese N° es de otro local: no se toca ni se dice de quién es
            continue
        # el área va en el local del empleado (si ya existía, el suyo; si es nuevo, el elegido)
        elid = cur["local_id"] if cur and cur["local_id"] is not None else lid
        dep_id = None
        if g(c_dep):
            dep = db.q1("SELECT id FROM departments WHERE name=? COLLATE NOCASE AND local_id=?", (g(c_dep), elid))
            dep_id = dep["id"] if dep else db.ex("INSERT INTO departments(name, local_id) VALUES (?,?)", (g(c_dep), elid))
        if cur:
            db.ex("UPDATE employees SET name=?, document=COALESCE(?,document), department_id=COALESCE(?,department_id), "
                  "position=COALESCE(?,position), email=COALESCE(?,email), phone=COALESCE(?,phone), "
                  "local_id=COALESCE(local_id, ?) WHERE id=?",
                  (name, g(c_doc), dep_id, g(c_pos), g(c_mail), g(c_tel), elid, cur["id"]))
        else:
            db.ex("INSERT INTO employees(employee_no,name,document,department_id,position,email,phone,schedule_id,local_id) "
                  "VALUES (?,?,?,?,?,?,?,?,?)", (no, name, g(c_doc), dep_id, g(c_pos), g(c_mail), g(c_tel), sched, lid))
        n += 1
    db.audit(u["username"], "importar excel", f"{n} en local {lid}" + (f"; omitidos de otro local: {', '.join(other)}" if other else ""))
    flash(request, f"{n} empleados importados/actualizados. Usa 'Sincronizar todos' para enviarlos a los equipos.")
    if other:
        flash(request, f"{len(other)} N° no se importaron porque ya los usa otro local: {', '.join(other[:20])}. "
                       f"El siguiente N° libre es {next_employee_no()}.", "warning")
    return back("/empleados")


@app.get("/fotos/{name}")
def photo(request: Request, name: str):
    u = user(request)
    name = Path(name).name
    # Solo la foto de un empleado que este usuario puede ver (el admin, cualquiera)
    if not alcance.is_admin(u):
        emp = db.q1("SELECT local_id FROM employees WHERE photo=?", (name,))
        if not emp or not alcance.can_see(request, u, emp["local_id"]):
            return Response(status_code=404)
    p = db.PHOTOS_DIR / name
    return FileResponse(p) if p.exists() else Response(status_code=404)


def _face_from_devices(emp):
    """Trae la foto de rostro registrada en el equipo; si no hay, la última captura de una marcación.
    Primero mira los equipos de su local."""
    devs = db.q("SELECT * FROM devices WHERE active=1 ORDER BY (local_id IS ?) DESC, id", (emp["local_id"],))
    for dev in devs:
        try:
            c = devices.client(dev)
            url = c.face_url(emp["employee_no"])
            if url:
                return c.fetch_picture(url), True
        except ISAPIError:
            pass
    ev = db.q1("""SELECT ev.picture_url, d.* FROM events ev JOIN devices d ON d.id=ev.device_id
                  WHERE ev.employee_no=? AND ev.picture_url IS NOT NULL AND ev.valid=1
                  ORDER BY ev.event_time DESC LIMIT 1""", (emp["employee_no"],))
    if ev:
        try:
            return devices.client(ev).fetch_picture(ev["picture_url"]), False
        except ISAPIError:
            pass
    return None, False


@app.get("/empleados/{eid}/rostro")
async def employee_face(request: Request, eid: int):
    """Foto del empleado: la del sistema o, si no tiene, la que está en el equipo (se guarda para la próxima)."""
    u = user(request)
    emp = alcance.row(request, u, "employees", eid)  # de otro local: como si no existiera
    if not emp:
        return Response(status_code=404)
    if emp["photo"] and (db.PHOTOS_DIR / emp["photo"]).exists():
        return FileResponse(db.PHOTOS_DIR / emp["photo"])
    data, enrolled = await run_in_threadpool(_face_from_devices, emp)
    if not data:
        return Response(status_code=404)
    if enrolled:
        name = f"{emp['employee_no']}_{uuid.uuid4().hex[:8]}.jpg"
        (db.PHOTOS_DIR / name).write_bytes(data)
        db.ex("UPDATE employees SET photo=? WHERE id=? AND photo IS NULL", (name, eid))
    return Response(data, media_type="image/jpeg", headers={"Cache-Control": "private, max-age=3600"})


# ------------------------------------------------------------------ departamentos / horarios / feriados
@app.get("/configuracion", response_class=HTMLResponse)
def config(request: Request):
    u = user(request, "admin", "encargado")
    cl, al = alcance.sql(request, u, "local_id", include_null=True)
    admin = alcance.is_admin(u)
    return render(request, "config.html", u, departments=db.q(f"SELECT * FROM departments WHERE {cl} ORDER BY name", al),
                  schedules=db.q(f"SELECT * FROM schedules WHERE {cl} ORDER BY name", al),
                  local_names={r["id"]: r["name"] for r in db.q("SELECT id, name FROM locales")},
                  new_local=alcance.default_local(request, u),
                  holidays=db.q("SELECT * FROM holidays ORDER BY date DESC") if admin else [],
                  settings={k: db.get_setting(k) for k in ("company", "poll_seconds", "min_gap_minutes")} if admin else {})


@app.post("/configuracion/{what}")
async def config_save(request: Request, what: str):
    u = user(request, "admin", "encargado")
    f = await form(request)

    def denied(msg):
        flash(request, msg, "danger")
        return back("/configuracion")

    # Feriados y ajustes generales valen para todos los locales: solo el admin
    if what in ("feriado", "feriado-eliminar", "general") and not alcance.is_admin(u):
        return denied("Solo el administrador puede cambiar feriados y ajustes generales.")
    # Áreas y horarios: se crean en el local elegido y solo se tocan los de un local permitido
    new_local = alcance.default_local(request, u)
    if (what == "departamento" or (what == "horario" and not f.get("id"))) and not new_local:
        return denied("Elige un local arriba en el menú.")
    if what.endswith("-eliminar") and not f.get("id"):
        return denied("Falta indicar qué eliminar.")
    if what == "departamento-eliminar" and not alcance.row(request, u, "departments", f.get("id")):
        return denied("Esa área no existe.")
    if what in ("horario", "horario-eliminar") and f.get("id") and not alcance.row(request, u, "schedules", f["id"]):
        return denied("Ese horario no existe.")
    if what == "departamento" and f.get("name"):
        if db.q1("SELECT 1 FROM departments WHERE local_id=? AND lower(trim(name))=lower(?)", (new_local, f["name"])):
            return denied(f"Ya existe el área «{f['name']}» en este local.")
        db.ex("INSERT INTO departments(name, local_id) VALUES (?,?)", (f["name"], new_local))
    elif what == "departamento-eliminar":
        db.ex("UPDATE employees SET department_id=NULL WHERE department_id=?", (f["id"],))
        db.ex("DELETE FROM departments WHERE id=?", (f["id"],))
    elif what == "horario":
        if not (f.get("name") and f.get("start_time") and f.get("end_time")):
            flash(request, "El horario necesita nombre, hora de entrada y hora de salida.", "danger")
            return back("/configuracion")
        if f["end_time"] <= f["start_time"]:
            flash(request, "La hora de salida debe ser después de la de entrada (los turnos de noche aún no se manejan).",
                  "danger")
            return back("/configuracion")
        if not all((f.get(k) or "0").isdigit() for k in ("tolerance_min", "lunch_min")):
            flash(request, "La tolerancia y el almuerzo deben ser minutos (números enteros).", "danger")
            return back("/configuracion")
        days = ",".join(sorted(v for k, v in (await request.form()).multi_items() if k == "days"))
        win = tuple(f.get(k) or None for k in ("in_from", "in_to", "out_from", "out_to"))
        if any(win) and not all(win):
            flash(request, "Completa las cuatro horas de las ventanas, o déjalas todas vacías para no usar ventanas.", "danger")
            return back("/configuracion")
        if all(win) and (win[0] > win[1] or win[2] > win[3]):
            flash(request, "En cada ventana la hora «desde» debe ser anterior a la hora «hasta».", "danger")
            return back("/configuracion")
        old = db.q1("SELECT * FROM schedules WHERE id=?", (f.get("id") or 0,))
        before = (old["start_time"], old["end_time"]) if old else ("08:00", "17:00")
        if all(win) and win == db.default_windows(*before) and before != (f["start_time"], f["end_time"]):
            # Las ventanas estaban en su valor automático: se mueven con las nuevas horas del horario
            win = db.default_windows(f["start_time"], f["end_time"])
        vals = (f["name"], f["start_time"], f["end_time"], int(f.get("tolerance_min") or 0), int(f.get("lunch_min") or 0),
                days) + win
        if f.get("id"):
            db.ex("""UPDATE schedules SET name=?, start_time=?, end_time=?, tolerance_min=?, lunch_min=?, days=?,
                     in_from=?, in_to=?, out_from=?, out_to=? WHERE id=?""", vals + (f["id"],))
        else:
            db.ex("""INSERT INTO schedules(name,start_time,end_time,tolerance_min,lunch_min,days,in_from,in_to,out_from,out_to,
                     local_id) VALUES (?,?,?,?,?,?,?,?,?,?,?)""", vals + (new_local,))
    elif what == "horario-eliminar":
        db.ex("UPDATE employees SET schedule_id=NULL WHERE schedule_id=?", (f["id"],))
        db.ex("DELETE FROM schedules WHERE id=?", (f["id"],))
    elif what == "feriado" and f.get("date"):
        db.ex("INSERT OR REPLACE INTO holidays(date, name) VALUES (?,?)", (f["date"], f.get("name")))
    elif what == "feriado-eliminar":
        db.ex("DELETE FROM holidays WHERE id=?", (f["id"],))
    elif what == "general" and u["role"] == "admin":
        limits = {"poll_seconds": (2, 300, "La revisión de respaldo debe ser de 2 a 300 segundos."),
                  "min_gap_minutes": (0, 60, "La separación mínima debe ser de 0 a 60 minutos.")}
        for k, (low, high, msg) in limits.items():
            if k in f and not (f[k].isdigit() and low <= int(f[k]) <= high):
                flash(request, msg, "danger")
                return back("/configuracion")
        for k in ("company", "poll_seconds", "min_gap_minutes"):
            if k in f:
                db.set_setting(k, f[k])
    db.audit(u["username"], f"configuración: {what}", str({k: v for k, v in f.items() if k != "password"}))
    flash(request, "Cambios guardados.")
    return back("/configuracion")


@app.get("/respaldo")
def backup_download(request: Request):
    u = user(request, "admin")
    data = db.backup_zip()
    db.audit(u["username"], "respaldo descargado", f"{len(data) // 1024} KB")
    name = f"respaldo-paso-largo-{datetime.now():%Y%m%d-%H%M}.zip"
    return Response(data, media_type="application/zip",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.post("/respaldo/restaurar")
async def backup_restore(request: Request):
    u = user(request, "admin")
    f = await request.form()
    up = f.get("archivo")
    if not up or not getattr(up, "filename", ""):
        flash(request, "Elige el archivo .zip del respaldo.", "danger")
        return back("/configuracion")
    raw = await up.read()
    # Copia de lo que hay ahora, para poder deshacer la restauración
    saved = db.DATA_DIR / "respaldos" / f"antes-de-restaurar-{datetime.now():%Y%m%d-%H%M%S}.zip"
    saved.parent.mkdir(exist_ok=True)
    saved.write_bytes(await run_in_threadpool(db.backup_zip))
    try:
        emps, evs = await run_in_threadpool(db.restore_zip, raw)
    except ValueError as e:
        flash(request, str(e), "danger")
        return back("/configuracion")
    db.audit(u["username"], "respaldo restaurado", f"{up.filename}: {emps} empleados, {evs} eventos; "
                                                   f"lo anterior quedó en datos/respaldos/{saved.name}")
    poller.wake.set()
    # Los usuarios ahora son los del respaldo: se pide entrar de nuevo
    request.session.clear()
    flash(request, f"Respaldo restaurado: {emps} empleados y {evs} marcaciones. Entra con un usuario del respaldo.")
    return back("/login")


# ------------------------------------------------------------------ permisos / justificaciones
@app.get("/permisos", response_class=HTMLResponse)
def leaves(request: Request):
    u = user(request, "admin", "encargado")
    ce, ae = alcance.sql(request, u, "e.local_id", include_null=True)
    rows = db.q(f"""SELECT l.*, e.name, e.employee_no FROM leaves l JOIN employees e ON e.id=l.employee_id
                    WHERE {ce} ORDER BY l.date_from DESC LIMIT 300""", ae)
    return render(request, "leaves.html", u, rows=rows, kinds=LEAVE_KINDS,
                  emps=db.q(f"SELECT id, name, employee_no FROM employees e WHERE active=1 AND {ce} ORDER BY name", ae))


@app.post("/permisos")
async def leave_save(request: Request):
    u = user(request, "admin", "encargado")
    f = await form(request)
    if f.get("delete"):
        lv = db.q1("SELECT * FROM leaves WHERE id=?", (f["delete"],))
        if not lv or not alcance.row(request, u, "employees", lv["employee_id"]):
            flash(request, "Ese permiso no existe.", "danger")
            return back("/permisos")
        db.ex("DELETE FROM leaves WHERE id=?", (f["delete"],))
    elif f.get("employee_id") and not alcance.row(request, u, "employees", f["employee_id"]):
        flash(request, "Ese empleado no existe.", "danger")
        return back("/permisos")
    elif f.get("employee_id") and f.get("date_from"):
        if f.get("kind") not in LEAVE_KINDS:
            flash(request, "Elige el tipo de permiso.", "danger")
            return back("/permisos")
        if (f.get("date_to") or f["date_from"]) < f["date_from"]:
            flash(request, "La fecha «hasta» no puede ser anterior a «desde».", "danger")
            return back("/permisos")
        db.ex("INSERT INTO leaves(employee_id,date_from,date_to,kind,notes,created_by) VALUES (?,?,?,?,?,?)",
              (f["employee_id"], f["date_from"], f.get("date_to") or f["date_from"], f.get("kind"), f.get("notes"),
               u["username"]))
    db.audit(u["username"], "permisos", str(f))
    flash(request, "Guardado.")
    return back("/permisos")


# ------------------------------------------------------------------ asistencia
@app.get("/asistencia", response_class=HTMLResponse)
def attendance_day(request: Request, fecha: str = "", dep: str = ""):
    u = user(request)
    d = attendance.parse_date(fecha, date.today())
    rows = attendance.compute(d, d, department_id=dep or None, locales=emp_locales(request, u))
    cl, al = alcance.sql(request, u, "local_id", include_null=True)
    return render(request, "attendance.html", u, rows=rows, d=d, dep=dep, prev=d - timedelta(days=1),
                  next=d + timedelta(days=1), departments=db.q(f"SELECT * FROM departments WHERE {cl} ORDER BY name", al),
                  emps=db.q(f"SELECT id, name, employee_no FROM employees WHERE active=1 AND {cl} ORDER BY name", al))


@app.post("/asistencia/manual")
async def manual_mark(request: Request):
    u = user(request, "admin", "encargado")
    f = await form(request)
    emp = alcance.row(request, u, "employees", f.get("employee_id"))
    if not emp:
        flash(request, "Ese empleado no existe.", "danger")
    elif f.get("date") and f.get("time"):
        try:
            t = datetime.strptime(f"{f['date']} {f['time']}", "%Y-%m-%d %H:%M").isoformat()
        except ValueError:
            flash(request, "Fecha u hora no válidas (usa el formato 07:45).", "danger")
            return back(f"/asistencia?fecha={f.get('date', '')}")
        db.ex("""INSERT INTO events(device_id, serial_no, event_time, employee_no, name, major, minor, verify_mode,
                 attendance_status, valid, source) VALUES (NULL, NULL, ?, ?, ?, 0, 0, 'manual', ?, 1, 'manual')""",
              (t, emp["employee_no"], emp["name"], f.get("kind") or None))
        db.audit(u["username"], "marca manual", f"{emp['employee_no']} {t} {f.get('kind')} {f.get('reason')}")
        flash(request, f"Marca manual registrada para {emp['name']}.")
    return back(f"/asistencia?fecha={f.get('date', '')}")


@app.get("/eventos", response_class=HTMLResponse)
def events(request: Request, desde: str = "", hasta: str = "", emp: str = "", dev: str = "", todos: str = ""):
    u = user(request)
    d1 = attendance.parse_date(desde, date.today())
    d2 = attendance.parse_date(hasta, d1)
    where, args = ["ev.event_time >= ?", "ev.event_time < ?"], [d1.isoformat(), (d2 + timedelta(days=1)).isoformat()]
    if emp:
        # solo se filtra por un empleado que este usuario puede ver
        where.append("ev.employee_no=?")
        args.append(emp if alcance.employee_by_no(request, u, emp) else None)
    if dev:
        where.append("ev.device_id=?")
        args.append(dev)
    if not todos:
        where.append("ev.employee_no IS NOT NULL")
    rows = visible_events(request, u, where, args, 1000, ", s.in_from, s.in_to, s.out_from, s.out_to",
                          "LEFT JOIN schedules s ON s.id=e.schedule_id")
    # también la gente de otros locales permitidos que marcó en estos equipos (si no, saldría «No cuenta»)
    roles = attendance.mark_roles(d1, d2, locales=alcance.allowed(u) + ([None] if alcance.is_admin(u) else []))
    for ev in rows:
        ev["counts_as"] = (roles.get((ev["employee_no"], datetime.fromisoformat(ev["event_time"])), "No cuenta")
                           if ev["valid"] and not ev.get("other") else "")
    cl, al = alcance.sql(request, u, "local_id", include_null=True)
    return render(request, "events.html", u, rows=rows, desde=d1, hasta=d2, emp=emp, dev=dev, todos=todos,
                  emps=db.q(f"SELECT name, employee_no FROM employees WHERE {cl} ORDER BY name", al),
                  devs=db.q(f"SELECT id, name FROM devices WHERE {cl} ORDER BY name", al))


def event_visible(request, u, ev) -> bool:
    """¿Puede este usuario ver quién hizo esta marca? Lo decide el local del empleado; si la marca no es de
    ningún empleado registrado, el local del equipo."""
    if alcance.is_admin(u):
        return True
    emp = db.q1("SELECT local_id FROM employees WHERE employee_no=?", (ev["employee_no"],)) if ev["employee_no"] else None
    if emp:
        return alcance.can_see(request, u, emp["local_id"])
    dev = db.q1("SELECT local_id FROM devices WHERE id=?", (ev["device_id"],)) if ev["device_id"] else None
    return bool(dev) and alcance.can_see(request, u, dev["local_id"])


@app.post("/eventos/{evid}/eliminar")
async def event_delete(request: Request, evid: int):
    """Solo se pueden borrar las marcas manuales; las del equipo son el registro original."""
    u = user(request, "admin", "encargado")
    ev = db.q1("SELECT * FROM events WHERE id=? AND source='manual'", (evid,))
    if ev and event_visible(request, u, ev):
        db.ex("DELETE FROM events WHERE id=?", (evid,))
        db.audit(u["username"], "borrar marca manual", f"{ev['employee_no']} {ev['event_time']} {ev['attendance_status']}")
        flash(request, "Marca manual borrada.")
    return back(seguridad.safe_back(request.headers.get("referer"), "/eventos"))


@app.get("/eventos/{evid}/foto")
async def event_photo(request: Request, evid: int):
    u = user(request)
    ev = db.q1("SELECT * FROM events WHERE id=?", (evid,))
    if ev and not event_visible(request, u, ev):
        return Response(status_code=404)
    dev = db.q1("SELECT * FROM devices WHERE id=?", (ev["device_id"],)) if ev else None
    if not dev or not ev["picture_url"]:
        return Response(status_code=404)
    try:
        data = await run_in_threadpool(devices.client(dev).fetch_picture, ev["picture_url"])
    except ISAPIError:
        return Response(status_code=502)
    # La foto de una marcación no cambia: el navegador la guarda y no se vuelve a pedir al equipo
    return Response(data, media_type="image/jpeg", headers={"Cache-Control": "private, max-age=604800"})


# ------------------------------------------------------------------ reportes
def report_args(request, u, desde, hasta, emp, dep):
    """Filas del reporte, solo de los empleados de los locales que se ven ahora."""
    today = date.today()
    d1 = attendance.parse_date(desde, today.replace(day=1))
    d2 = attendance.parse_date(hasta, today)
    rows = attendance.compute(d1, d2, employee_id=emp or None, department_id=dep or None,
                              locales=emp_locales(request, u))
    return d1, d2, rows, attendance.summarize(rows)


@app.get("/reportes", response_class=HTMLResponse)
def report(request: Request, desde: str = "", hasta: str = "", emp: str = "", dep: str = "", vista: str = "resumen"):
    u = user(request)
    d1, d2, rows, summary = report_args(request, u, desde, hasta, emp, dep)
    cl, al = alcance.sql(request, u, "local_id", include_null=True)
    return render(request, "reports.html", u, d1=d1, d2=d2, emp=emp, dep=dep, vista=vista, rows=rows,
                  summary=summary, departments=db.q(f"SELECT * FROM departments WHERE {cl} ORDER BY name", al),
                  emps=db.q(f"SELECT id, name FROM employees WHERE {cl} ORDER BY name", al), qs=request.url.query)


@app.get("/reportes/exportar.{fmt}")
def report_export(request: Request, fmt: str, desde: str = "", hasta: str = "", emp: str = "", dep: str = ""):
    u = user(request)
    d1, d2, rows, summary = report_args(request, u, desde, hasta, emp, dep)
    loc = alcance.current(request, u)
    title = "Reporte de asistencia" + (f" — {loc['name']}" if loc else "")
    period = f"Del {d1.strftime('%d/%m/%Y')} al {d2.strftime('%d/%m/%Y')}"
    fname = f"asistencia_{d1.isoformat()}_{d2.isoformat()}"
    db.audit(u["username"], f"exportar {fmt}", period)
    if fmt == "xlsx":
        return Response(reports.to_excel(rows, summary, title, period),
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="{fname}.xlsx"'})
    return Response(reports.to_pdf(rows, summary, title, period, db.get_setting("company", "")),
                    media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{fname}.pdf"'})


# ------------------------------------------------------------------ dispositivos
def _devices_page(request, u, **ctx):
    """Página de equipos: solo los de los locales que se ven ahora. La clave del equipo nunca va a la página,
    y el usuario del equipo solo a quien puede editarlo."""
    cond, args = alcance.sql(request, u, "d.local_id", include_null=True)
    devs = db.q(f"""SELECT d.*, l.name AS local, (SELECT COUNT(*) FROM events ev WHERE ev.device_id=d.id) AS n_events
                    FROM devices d LEFT JOIN locales l ON l.id=d.local_id WHERE {cond} ORDER BY l.name, d.name""", args)
    can_edit = u["role"] in alcance.EDIT_ROLES
    for d in devs:
        d.pop("password", None)
        if not can_edit:
            d.pop("username", None)
    # «Revisar empleados entre equipos»: solo entre equipos de un mismo local con 2 o más activos
    per_local = defaultdict(list)
    for d in devs:
        if d["active"] and d["local_id"] is not None:
            per_local[(d["local_id"], d["local"])].append(d)
    compare_locals = [{"id": k[0], "name": k[1]} for k, v in per_local.items() if len(v) > 1]
    return render(request, "devices.html", u, devs=devs, can_edit=can_edit, compare_locals=compare_locals,
                  local_opts=alcance.choices(u) if can_edit else [],
                  default_local=alcance.default_local(request, u), **ctx)


def _addr(host):
    """IP real de un host («localhost» y nombres se resuelven), para no confundir dos formas de la misma dirección."""
    import socket
    try:
        return socket.gethostbyname(host.strip().lower())
    except (OSError, UnicodeError):
        return host.strip().lower()


def _foreign_at(request, u, host, port, exclude_id=0):
    """¿Hay un equipo de otro local (que este usuario no ve) en esa misma dirección y puerto?"""
    ip = _addr(host)
    return any(not alcance.can_see(request, u, d["local_id"]) and _addr(d["host"]) == ip
               for d in db.q("SELECT host, local_id FROM devices WHERE port=? AND id<>?", (port, exclude_id)))


def _device_visible(request, u, did):
    """El equipo si este usuario puede verlo (su local); si no, None (se trata como inexistente)."""
    return alcance.row(request, u, "devices", did)


@app.get("/dispositivos", response_class=HTMLResponse)
def device_list(request: Request):
    u = user(request)
    return _devices_page(request, u)


@app.post("/dispositivos/buscar")
async def device_search(request: Request):
    """Busca biométricos en la red; a los ya registrados que cambiaron de IP se les corrige sola."""
    u = user(request, "admin")
    found, nets = await run_in_threadpool(devices.search_network)
    for f in found:
        if f["status"] == "movido":
            flash(request, f"«{f['dev_name']}» estaba en otra IP: {f['note']}. Ya se reconectó.")
    if not found:
        flash(request, f"No se encontró ningún biométrico en la red ({', '.join(nets) or 'sin red'}). "
                       "Revisa que esté encendido y conectado al mismo router/switch que esta PC.", "warning")
    db.audit(u["username"], "buscar equipos en la red", ", ".join(f["ip"] for f in found))
    poller.wake.set()
    return _devices_page(request, u, found=found, nets=nets, searched=True)


def _compare_local(request, u, value):
    """Local cuyos equipos se comparan: el indicado (si es suyo) o el que se está viendo."""
    lid = int(value) if str(value or "").isdigit() else alcance.default_local(request, u)
    return lid if lid is not None and lid in alcance.allowed(u) else None


@app.get("/dispositivos/comparar", response_class=HTMLResponse)
async def device_compare(request: Request, local: str = ""):
    """Con dos o más biométricos en un local: revisa que cada persona tenga el mismo N° y su rostro en todos."""
    u = user(request, "admin", "encargado")
    lid = _compare_local(request, u, local)
    if lid is None:
        flash(request, "Elige de qué local revisar los equipos.", "warning")
        return back("/dispositivos")
    ids = [d["id"] for d in devices.local_devices(lid)]
    cdevs, rows, errors = await run_in_threadpool(devices.compare_users, ids)
    names = {d["id"]: d["name"] for d in devices.local_devices(lid)}
    return _devices_page(request, u, compared=True, compare_local=lid, cdevs=cdevs, crows=rows,
                         cerrors={names.get(k, k): v for k, v in errors.items()})


@app.post("/dispositivos/copiar")
async def device_copy(request: Request):
    u = user(request, "admin", "encargado")
    raw = await request.form()
    lid = _compare_local(request, u, raw.get("local"))
    if lid is None:
        flash(request, "Elige de qué local copiar los rostros.", "warning")
        return back("/dispositivos")
    nos = [v for k, v in raw.multi_items() if k == "no"]
    ids = [d["id"] for d in devices.local_devices(lid)]
    done, msgs = await run_in_threadpool(devices.copy_to_missing, nos, ids)
    db.audit(u["username"], "copiar rostros entre equipos", f"local {lid}: " + ", ".join(nos))
    if done:
        flash(request, f"Se copiaron {done} rostros a los equipos donde faltaban.")
    for m in msgs[:8]:
        flash(request, m, "warning")
    return back(f"/dispositivos/comparar?local={lid}")


@app.get("/dispositivos/{did}/rostro/{no}")
async def device_face(request: Request, did: int, no: str):
    """Foto del rostro guardado en ese equipo para ese N° (para comparar entre equipos)."""
    u = user(request, "admin", "encargado")
    dev = _device_visible(request, u, did)
    if not dev or dev.get("paused_other"):
        return Response(status_code=404)

    def get():
        c = devices.client(dev)
        url = c.face_url(no)
        return c.fetch_picture(url) if url else None
    try:
        data = await run_in_threadpool(get)
    except ISAPIError:
        data = None
    if not data:
        return Response(status_code=404)
    return Response(data, media_type="image/jpeg", headers={"Cache-Control": "private, max-age=300"})


@app.post("/dispositivos/guardar")
async def device_save(request: Request):
    u = user(request, "admin", "encargado")
    f = await form(request)
    if not (f.get("port") or "80").isdigit() or not f.get("host") or not f.get("name") or not f.get("username"):
        flash(request, "Revisa los datos del equipo: nombre, IP, puerto (número) y usuario son obligatorios.", "danger")
        return back("/dispositivos")
    old = None
    if f.get("id"):
        old = _device_visible(request, u, f["id"]) if f["id"].isdigit() else None
        if not old:
            flash(request, "Ese equipo no existe o es de otro local.", "danger")
            return back("/dispositivos")
    # Local del equipo: el elegido (solo uno de los suyos), o el que ya tenía, o el que se está viendo
    lid = int(f["local_id"]) if (f.get("local_id") or "").isdigit() else (
        old["local_id"] if old else alcance.default_local(request, u))
    if lid is None or lid not in alcance.allowed(u):
        flash(request, "Elige el local donde está el equipo (solo puedes usar tus locales).", "danger")
        return back("/dispositivos")
    host, port = f["host"].strip().lower(), int(f.get("port") or 80)
    vals = [f["name"], host, port, 1 if f.get("https") else 0, f["username"],
            devices.local_offset(), int(f.get("door_no") or 1), f.get("direction") or "ambos",
            1 if f.get("active") else 0, lid]
    # Un equipo se identifica por IP + puerto: nunca se registra dos veces
    same = db.q1("SELECT id, name, local_id FROM devices WHERE host=? AND port=? AND id<>?",
                 (host, port, old["id"] if old else 0))
    if (same and not alcance.can_see(request, u, same["local_id"])) or \
            await run_in_threadpool(_foreign_at, request, u, host, port, old["id"] if old else 0):
        # es de otro local: no se toca ni se dice cuál es
        flash(request, f"La dirección {host}:{port} ya está registrada en otro local. Pídele al administrador que lo revise.",
              "danger")
        return back("/dispositivos")
    if same and old:
        flash(request, f"La dirección {host}:{port} ya es del equipo «{same['name']}».", "danger")
        return back("/dispositivos")
    if same:
        old = same
        vals[-2] = 1
        flash(request, f"El equipo {host}:{port} ya estaba registrado como «{same['name']}»: se actualizaron "
                       "sus datos en lugar de crear otro.", "info")
    if old and not alcance.is_admin(u) and db.q1("SELECT paused_other FROM devices WHERE id=?", (old["id"],))["paused_other"]:
        vals[-2] = 0  # pausado por ser de otro local: sigue pausado hasta que el admin lo revise
    if old:
        db.ex("""UPDATE devices SET name=?, host=?, port=?, https=?, username=?, tz_offset=?, door_no=?, direction=?,
                 active=?, local_id=? WHERE id=?""", vals + [old["id"]])
        if f.get("password"):
            db.ex("UPDATE devices SET password=? WHERE id=?", (f["password"], old["id"]))
        did = old["id"]
    else:
        did = db.ex("""INSERT INTO devices(name, host, port, https, username, tz_offset, door_no, direction, active,
                       local_id, password) VALUES (?,?,?,?,?,?,?,?,?,?,?)""", vals + [f.get("password", "")])
    db.audit(u["username"], "guardar dispositivo", f"{f['name']} {f['host']} (local {lid})")
    dev = db.q1("SELECT * FROM devices WHERE id=?", (did,))
    try:  # primero se revisa que no sea el aparato de otro local; recién entonces lo lee el lector
        info = await run_in_threadpool(devices.refresh_info, dev)
        poller.wake.set()
        flash(request, f"Conectado a {info.get('deviceName') or dev['name']} — modelo {info.get('model')}, "
                       f"serie {info.get('serialNumber')}, firmware {info.get('firmwareVersion')}.")
    except ISAPIError as e:
        flash(request, f"Dispositivo guardado, pero no responde: {e}", "warning")
    return back("/dispositivos")


@app.post("/dispositivos/{did}/{action}")
async def device_action(request: Request, did: int, action: str):
    u = user(request, "admin", "encargado")
    dev = _device_visible(request, u, did)
    if not dev:
        flash(request, "Ese equipo no existe o es de otro local.", "danger")
        return back("/dispositivos")
    if dev.get("paused_other") and action != "eliminar" and not (alcance.is_admin(u) and action in ("activar", "ip")):
        # es el aparato de otro local: no se le habla (ni se bajan sus marcas) hasta que el admin lo revise
        flash(request, "Ese equipo se pausó porque ya está registrado en otro local: solo el administrador "
                       "puede reactivarlo o cambiarle la IP.", "danger")
        return back("/dispositivos")
    c = devices.client(dev)
    try:
        if action == "ip" and not alcance.is_admin(u):
            # cambiar a la IP de un equipo de otro local lo uniría con ese registro: solo el admin
            f = await form(request)
            host, _, port = (f.get("host") or "").strip().partition(":")
            if _foreign_at(request, u, host, int(port) if port.isdigit() else dev["port"], did):
                flash(request, f"La IP {host} ya está registrada en otro local. Pídele al administrador que lo revise.",
                      "danger")
                return back("/dispositivos")
        if action == "probar":
            info = await run_in_threadpool(devices.refresh_info, dev)
            n = await run_in_threadpool(c.user_count)
            t = await run_in_threadpool(c.get_time)
            flash(request, f"En línea: {info.get('model')} (serie {info.get('serialNumber')}), {n} usuarios, "
                           f"hora del equipo {t.get('localTime', '?')}.")
        elif action == "ip":
            f = await form(request)
            host, _, port = (f.get("host") or "").strip().partition(":")
            if not re.fullmatch(r"[\w.-]+", host):
                flash(request, "Escribe la IP del equipo, por ejemplo 192.168.1.200.", "danger")
                return back("/dispositivos")
            info = await run_in_threadpool(devices.connect_at, dev, host, port if port.isdigit() else None)
            poller.wake.set()
            flash(request, f"Conectado: {dev['name']} ahora está en {host.lower()} "
                           f"({info.get('model')}, serie {info.get('serialNumber')}).")
        elif action == "hora":
            await run_in_threadpool(c.set_time, datetime.now(), devices.local_offset())
            devices.CLOCK_WARNINGS.pop(did, None)
            flash(request, f"Hora de {dev['name']} sincronizada con este servidor ({datetime.now():%d/%m/%Y %H:%M}).")
        elif action == "eventos":
            dias = (await form(request)).get("dias") or "31"
            since = datetime.now() - timedelta(days=int(dias) if dias.isdigit() else 31)
            n = await run_in_threadpool(devices.pull_events, dev, since)
            flash(request, f"Descarga completa: {n} eventos nuevos.")
        elif action == "reiniciar":
            await run_in_threadpool(c.reboot)
            flash(request, f"{dev['name']} se está reiniciando.", "warning")
        elif action == "eliminar":
            db.ex("UPDATE devices SET active=0 WHERE id=?", (did,))
            devices.CLOCK_WARNINGS.pop(did, None)
            flash(request, "Dispositivo desactivado: ya no se leen sus marcas (su historial se conserva).", "warning")
        elif action == "activar":
            db.ex("UPDATE devices SET active=1 WHERE id=?", (did,))
            try:  # se revisa enseguida que no sea el aparato de otro local (si lo es, vuelve a quedar pausado)
                await run_in_threadpool(devices.refresh_info, db.q1("SELECT * FROM devices WHERE id=?", (did,)))
            except ISAPIError as e:
                if devices.OTHER_LOCAL in str(e):
                    raise
            poller.wake.set()
            flash(request, f"{dev['name']} vuelve a estar activo.")
        db.audit(u["username"], f"dispositivo {action}", dev["name"])
    except ISAPIError as e:
        flash(request, f"{dev['name']}: {e}", "danger")
    return back(seguridad.safe_back(request.headers.get("referer"), "/dispositivos"))


# ------------------------------------------------------------------ usuarios del sistema
@app.get("/usuarios", response_class=HTMLResponse)
def users(request: Request):
    u = user(request, "admin")
    rows = db.q("SELECT * FROM users ORDER BY username")
    mine = defaultdict(list)
    for r in db.q("SELECT ul.user_id, ul.local_id FROM user_locales ul JOIN locales l ON l.id=ul.local_id WHERE l.active=1"):
        mine[r["user_id"]].append(r["local_id"])
    for r in rows:
        r["locales"] = mine[r["id"]]
        r["sin_local"] = r["role"] != "admin" and not r["locales"]
    return render(request, "users.html", u, users=rows, locales=db.q("SELECT * FROM locales WHERE active=1 ORDER BY name"),
                  log=db.q("SELECT * FROM audit ORDER BY id DESC LIMIT 200"))


def user_locales_from(fd):
    """Locales activos marcados en el formulario de usuario (casillas name=locales)."""
    ok = {r["id"] for r in db.q("SELECT id FROM locales WHERE active=1")}
    out = set()
    for v in fd.getlist("locales"):
        try:
            if int(v) in ok:
                out.add(int(v))
        except ValueError:
            pass
    return sorted(out)


def set_user_locales(uid, role, lids):
    """Guarda los locales del usuario. El admin ve todos: no se le guarda ninguno.
    Los locales inactivos que ya tenía se conservan (por si se vuelven a activar)."""
    with db.tx() as c:
        c.execute("DELETE FROM user_locales WHERE user_id=? AND local_id IN (SELECT id FROM locales WHERE active=1)", (uid,))
        if role == "admin":
            c.execute("DELETE FROM user_locales WHERE user_id=?", (uid,))
        else:
            c.executemany("INSERT OR IGNORE INTO user_locales(user_id, local_id) VALUES (?,?)", [(uid, lid) for lid in lids])


@app.post("/usuarios")
async def user_save(request: Request):
    u = user(request, "admin")
    f = await form(request)
    lids = user_locales_from(await request.form())
    f.setdefault("role", "consulta")
    if f["role"] not in ROLES:
        flash(request, "Rol no válido.", "danger")
        return back("/usuarios")
    if f["role"] != "admin" and not lids:
        flash(request, f"Marca al menos un local para el usuario con rol «{ROLES[f['role']]}».", "danger")
        return back("/usuarios")
    if f.get("id"):
        if not db.q1("SELECT 1 FROM users WHERE id=?", (f["id"],)):
            flash(request, "Ese usuario ya no existe.", "danger")
            return back("/usuarios")
        if int(f["id"]) == u["id"] and (f.get("role") != "admin" or not f.get("active")):
            flash(request, "No puedes quitarte el rol de administrador ni desactivarte.", "danger")
            return back("/usuarios")
        db.ex("UPDATE users SET full_name=?, role=?, active=? WHERE id=?",
              (f.get("full_name"), f["role"], 1 if f.get("active") else 0, f["id"]))
        if f.get("password"):
            if len(f["password"]) < MIN_PASSWORD:
                flash(request, f"La contraseña debe tener {MIN_PASSWORD}+ caracteres.", "danger")
                return back("/usuarios")
            new_hash = db.hash_password(f["password"])
            db.ex("UPDATE users SET password_hash=? WHERE id=?", (new_hash, f["id"]))
            if int(f["id"]) == u["id"]:
                request.session["pv"] = seguridad.password_tag(SECRET, new_hash)
    else:
        if not f.get("username") or len(f.get("password", "")) < MIN_PASSWORD:
            flash(request, f"Usuario y contraseña ({MIN_PASSWORD}+ caracteres) son obligatorios.", "danger")
            return back("/usuarios")
        if db.q1("SELECT 1 FROM users WHERE username=?", (f["username"],)):
            flash(request, "Ese nombre de usuario ya existe.", "danger")
            return back("/usuarios")
        f["id"] = db.ex("INSERT INTO users(username, full_name, password_hash, role) VALUES (?,?,?,?)",
                        (f["username"], f.get("full_name"), db.hash_password(f["password"]), f["role"]))
    set_user_locales(int(f["id"]), f["role"], lids)
    names = [r["name"] for r in db.q(f"SELECT name FROM locales WHERE id IN ({','.join('?' * len(lids))}) ORDER BY name", lids)] if lids else []
    who = f.get("username") or (db.q1("SELECT username FROM users WHERE id=?", (f["id"],)) or {}).get("username")
    db.audit(u["username"], "guardar usuario",
             f"{who} · {ROLES[f['role']]}" + ("" if f["role"] == "admin" else f" · locales: {', '.join(names)}"))
    flash(request, "Usuario guardado.")
    return back("/usuarios")


# ------------------------------------------------------------------ locales (solo admin)
NULL_LABELS = {"devices": "biométricos", "departments": "áreas", "employees": "empleados", "schedules": "horarios"}


def users_without_local():
    """Usuarios activos que no son admin y no tienen ningún local activo: no pueden ver nada."""
    return db.q("""SELECT username FROM users u WHERE u.active=1 AND u.role!='admin' AND NOT EXISTS (
                       SELECT 1 FROM user_locales ul JOIN locales l ON l.id=ul.local_id WHERE ul.user_id=u.id AND l.active=1)
                   ORDER BY username""")


@app.get("/locales", response_class=HTMLResponse)
def locales_page(request: Request):
    u = user(request, "admin")
    rows = db.q("""SELECT l.*,
        (SELECT COUNT(*) FROM devices d WHERE d.local_id=l.id) dev_total,
        (SELECT COUNT(*) FROM devices d WHERE d.local_id=l.id AND d.active=1 AND d.online=1) dev_online,
        (SELECT COUNT(*) FROM employees e WHERE e.local_id=l.id AND e.active=1) emps,
        (SELECT COUNT(*) FROM departments a WHERE a.local_id=l.id) areas,
        (SELECT COUNT(*) FROM user_locales ul JOIN users x ON x.id=ul.user_id
          WHERE ul.local_id=l.id AND x.active=1 AND x.role!='admin') users
        FROM locales l ORDER BY l.active DESC, l.name""")
    nulls = {NULL_LABELS[t]: n for t in db.LOCAL_TABLES
             if (n := db.q1(f"SELECT COUNT(*) n FROM {t} WHERE local_id IS NULL")["n"])}
    return render(request, "locales.html", u, locales=rows, nulls=nulls, sin_local=users_without_local(),
                  devs=db.q("SELECT id, name, host, port, local_id, online, active FROM devices ORDER BY name"),
                  areas=db.q("""SELECT a.id, a.name, a.local_id,
                                (SELECT COUNT(*) FROM employees e WHERE e.department_id=a.id) emps
                                FROM departments a ORDER BY a.name"""),
                  emps=db.q("""SELECT e.id, e.employee_no, e.name, e.local_id, e.active, a.name area
                               FROM employees e LEFT JOIN departments a ON a.id=e.department_id
                               ORDER BY e.active DESC, e.name"""))


@app.post("/locales")
async def locales_save(request: Request):
    u, f = user(request, "admin"), await form(request)
    action, name = f.get("action"), " ".join(f.get("name", "").split())[:60]
    loc = db.q1("SELECT * FROM locales WHERE id=?", (f.get("id"),)) if f.get("id") else None
    if action in ("crear", "renombrar"):
        if not name:
            flash(request, "Escribe el nombre del local.", "danger")
        elif db.q1("SELECT 1 FROM locales WHERE name=? COLLATE NOCASE AND id!=?", (name, loc["id"] if loc else 0)):
            flash(request, f"Ya existe un local llamado «{name}».", "danger")
        elif action == "crear":
            lid = db.ex("INSERT INTO locales(name) VALUES (?)", (name,))
            # Un horario de partida, para que los empleados que lleguen desde sus biométricos tengan uno
            db.ex("""INSERT INTO schedules(name,start_time,end_time,tolerance_min,lunch_min,days,in_from,in_to,out_from,out_to,
                     local_id) VALUES ('Oficina L-V','08:00','17:00',10,60,'1,2,3,4,5',?,?,?,?,?)""",
                  db.default_windows("08:00", "17:00") + (lid,))
            db.audit(u["username"], "crear local", name)
            flash(request, f"Local «{name}» creado. Ahora puedes moverle biométricos, áreas y empleados abajo.")
        elif loc:
            db.ex("UPDATE locales SET name=? WHERE id=?", (name, loc["id"]))
            db.audit(u["username"], "renombrar local", f"{loc['name']} → {name}")
            flash(request, f"Local renombrado a «{name}».")
    elif action in ("activar", "desactivar") and loc:
        if action == "desactivar":
            if loc["active"] and db.q1("SELECT COUNT(*) n FROM locales WHERE active=1")["n"] <= 1:
                flash(request, "No se puede desactivar el último local activo.", "danger")
                return back("/locales")
            db.ex("UPDATE locales SET active=0 WHERE id=?", (loc["id"],))
            flash(request, f"Local «{loc['name']}» desactivado: nadie lo ve en el menú; sus datos se conservan.", "warning")
        else:
            db.ex("UPDATE locales SET active=1 WHERE id=?", (loc["id"],))
            flash(request, f"Local «{loc['name']}» activado.")
        db.audit(u["username"], f"{action} local", loc["name"])
        if left := users_without_local():
            flash(request, "Quedan usuarios sin ningún local activo: " + ", ".join(x["username"] for x in left)
                  + ". Asígnales uno en Usuarios.", "warning")
    else:
        flash(request, "Acción no válida.", "danger")
    return back("/locales")


async def _move_on_devices(request, eids, old_locals):
    """Quien cambia de local sale de los biométricos de su local anterior y entra a los del nuevo."""
    errors = 0
    for eid in eids:
        emp = db.q1("SELECT * FROM employees WHERE id=?", (eid,))
        if not emp or not emp["active"]:
            continue
        old = old_locals.get(eid)
        if old and old != emp["local_id"]:
            errors += sum(not ok for _, ok, _m in await run_in_threadpool(devices.remove_employee, emp, [old]))
        errors += sum(not ok for _, ok, _m in await run_in_threadpool(devices.push_employee, emp))
    if errors:
        flash(request, f"{errors} envío(s) a biométricos fallaron (¿equipo apagado?). Usa «Sincronizar» en Empleados "
                       "cuando estén en línea.", "warning")


@app.post("/locales/mover")
async def locales_move(request: Request):
    """Revisar reparto: pasa un biométrico, un área (con su gente) o varios empleados a otro local."""
    u, f = user(request, "admin"), await form(request)
    what = f.get("que")
    dest = db.q1("SELECT * FROM locales WHERE id=?", (f.get("destino") or 0,))
    if not dest:
        flash(request, "Elige el local de destino.", "danger")
        return back("/locales#reparto")
    lname = lambda lid: (db.q1("SELECT name FROM locales WHERE id=?", (lid,)) or {"name": "sin local"})["name"]
    if what == "dispositivo":
        dev = db.q1("SELECT * FROM devices WHERE id=?", (f.get("id") or 0,))
        if not dev:
            flash(request, "Ese biométrico ya no existe.", "danger")
        elif dev["local_id"] != dest["id"]:
            db.ex("UPDATE devices SET local_id=? WHERE id=?", (dest["id"], dev["id"]))
            db.audit(u["username"], "mover biométrico de local", f"{dev['name']}: {lname(dev['local_id'])} → {dest['name']}")
            flash(request, f"Biométrico «{dev['name']}» ahora es de {dest['name']}.")
            if dev["active"]:
                dev = db.q1("SELECT * FROM devices WHERE id=?", (dev["id"],))
                errors = 0
                for emp in db.q("SELECT * FROM employees WHERE active=1 AND local_id=?", (dest["id"],)):
                    errors += sum(not ok for _, ok, _m in await run_in_threadpool(devices.push_employee, emp, [dev]))
                if errors:
                    flash(request, f"{errors} empleado(s) de {dest['name']} no se pudieron enviar al equipo (¿apagado?). "
                                   "Usa «Sincronizar todos» en Empleados cuando esté en línea.", "warning")
                flash(request, "La gente del local anterior que estaba en ese equipo sigue registrada en él; si ya no "
                               "debe marcar ahí, bórrala desde Dispositivos → Revisar empleados entre equipos.", "info")
        return back("/locales#reparto")
    if what == "area":
        area = db.q1("SELECT * FROM departments WHERE id=?", (f.get("id") or 0,))
        if not area:
            flash(request, "Esa área ya no existe.", "danger")
            return back("/locales#reparto")
        if area["local_id"] == dest["id"]:
            return back("/locales#reparto")
        if db.q1("SELECT 1 FROM departments WHERE local_id=? AND name=? COLLATE NOCASE AND id!=?",
                 (dest["id"], area["name"], area["id"])):
            flash(request, f"{dest['name']} ya tiene un área llamada «{area['name']}». Cambia el nombre de una de las dos antes de moverla.", "danger")
            return back("/locales#reparto")
        moved = {r["id"]: r["local_id"] for r in db.q("SELECT id, local_id FROM employees WHERE department_id=?", (area["id"],))}
        with db.tx() as c:
            c.execute("UPDATE departments SET local_id=? WHERE id=?", (dest["id"], area["id"]))
            n = c.execute("UPDATE employees SET local_id=? WHERE department_id=?", (dest["id"], area["id"])).rowcount
            # el horario es de cada local: si el que tenían no es del nuevo local, se les quita
            lost = c.execute("""UPDATE employees SET schedule_id=NULL WHERE department_id=? AND schedule_id IS NOT NULL
                                AND schedule_id NOT IN (SELECT id FROM schedules WHERE local_id=?)""",
                             (area["id"], dest["id"])).rowcount
        db.audit(u["username"], "mover área de local",
                 f"{area['name']} ({n} empleados): {lname(area['local_id'])} → {dest['name']}" + (f"; {lost} sin horario" if lost else ""))
        flash(request, f"Área «{area['name']}» y sus {n} empleados ahora son de {dest['name']}.")
        await _move_on_devices(request, list(moved), moved)
        if lost:
            flash(request, f"{lost} de esos empleados quedaron sin horario (el suyo era de otro local). "
                           f"Asígnales un horario de {dest['name']} en Empleados.", "warning")
        return back("/locales#reparto")
    if what == "empleados":
        eids = []
        for v in (await request.form()).getlist("ids"):
            try:
                eids.append(int(v))
            except ValueError:
                pass
        rows = db.q(f"SELECT id, name, local_id FROM employees WHERE id IN ({','.join('?' * len(eids))})", eids) if eids else []
        rows = [r for r in rows if r["local_id"] != dest["id"]]
        if not rows:
            flash(request, "Marca al menos un empleado de otro local para mover.", "warning")
            return back("/locales#reparto")
        ids_ = [r["id"] for r in rows]
        marks = ",".join("?" * len(ids_))
        with db.tx() as c:
            c.execute(f"UPDATE employees SET local_id=? WHERE id IN ({marks})", [dest["id"]] + ids_)
            no_area = c.execute(f"""UPDATE employees SET department_id=NULL WHERE id IN ({marks}) AND department_id IS NOT NULL
                                    AND department_id NOT IN (SELECT id FROM departments WHERE local_id=?)""",
                                ids_ + [dest["id"]]).rowcount
            no_sched = c.execute(f"""UPDATE employees SET schedule_id=NULL WHERE id IN ({marks}) AND schedule_id IS NOT NULL
                                     AND schedule_id NOT IN (SELECT id FROM schedules WHERE local_id=?)""",
                                 ids_ + [dest["id"]]).rowcount
            pending = c.execute(f"SELECT COUNT(*) FROM employees WHERE id IN ({marks}) AND schedule_id IS NULL", ids_).fetchone()[0]
        await _move_on_devices(request, ids_, {r["id"]: r["local_id"] for r in rows})
        names = ", ".join(r["name"] for r in rows[:20]) + (f" y {len(rows) - 20} más" if len(rows) > 20 else "")
        db.audit(u["username"], "mover empleados de local", f"{len(rows)} → {dest['name']}: {names}")
        flash(request, f"{len(rows)} empleado(s) movidos a {dest['name']}."
                       + (f" A {no_area} se les quitó el área (era de otro local)." if no_area else ""))
        if pending:
            flash(request, f"{pending} empleado(s) movidos no tienen horario"
                           + (f" ({no_sched} lo perdieron por ser de otro local)" if no_sched else "")
                           + f". Asígnales un horario de {dest['name']} en Empleados.", "warning")
        return back("/locales#reparto")
    flash(request, "Acción no válida.", "danger")
    return back("/locales")


# ------------------------------------------------------------------ resumen de locales (solo admin)
from . import resumen  # noqa: E402


@app.get("/resumen", response_class=HTMLResponse)
def locales_summary(request: Request, vista: str = "hoy"):
    u = user(request, "admin")
    return render(request, "resumen.html", u, r=resumen.build(vista))


@app.get("/resumen/exportar.xlsx")
def locales_summary_export(request: Request, vista: str = "hoy"):
    u = user(request, "admin")
    data = resumen.build(vista)
    db.audit(u["username"], "exportar resumen de locales", f"{data['title']}: {resumen.period_text(data)}")
    fname = f"resumen_locales_{data['vista']}_{data['end'].isoformat()}"
    return Response(resumen.to_excel(data, db.get_setting("company", "")),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{fname}.xlsx"'})
