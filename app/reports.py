"""Exportación de reportes de asistencia a Excel y PDF."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .attendance import hm

DAYS = ["", "Lun", "Mar", "Mié", "Jue", "Vie", "Sáb", "Dom"]
DETAIL_HEAD = ["Fecha", "Día", "N°", "Empleado", "Área", "Entrada", "Salida",
               "Tardanza", "Horas", "Extra", "Estado", "Observación"]
SUMMARY_HEAD = ["N°", "Empleado", "Área", "Días asistidos", "Tardanzas", "Min. tarde",
                "Faltas", "Justificadas", "Sin salida", "Horas trabajadas", "Horas extra"]
FILLS = {"Tarde": "FFF3CD", "Ausente": "F8D7DA", "Justificado": "D1ECF1", "Feriado": "E2E3E5"}


def _detail(r):
    return [r["date"].strftime("%d/%m/%Y"), DAYS[r["weekday"]], r["employee_no"], r["name"], r["department"],
            r["entry"].strftime("%H:%M") if r["entry"] else "", r["exit"].strftime("%H:%M") if r["exit"] else "",
            r["late_min"] or "", hm(r["worked_min"]) if r["worked_min"] else "",
            hm(r["overtime_min"]) if r["overtime_min"] else "", r["status"], r["note"]]


def _summary(s):
    return [s["employee_no"], s["name"], s["department"], s["days_present"], s["late_days"], s["late_min"],
            s["absences"], s["justified"], s["missing_exit"], hm(s["worked_min"]), hm(s["overtime_min"])]


def to_excel(rows, summary, title, period):
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79")

    def sheet(ws, head, data, status_col=None):
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([period])
        ws.append([])
        ws.append(head)
        for cell in ws[4]:
            cell.font, cell.fill, cell.alignment = head_font, head_fill, Alignment(horizontal="center")
        for line in data:
            ws.append(line)
            if status_col is not None and line[status_col] in FILLS:
                for cell in ws[ws.max_row]:
                    cell.fill = PatternFill("solid", fgColor=FILLS[line[status_col]])
        for i, _ in enumerate(head, 1):
            width = max(len(str(ws.cell(row=r, column=i).value or "")) for r in range(4, ws.max_row + 1))
            ws.column_dimensions[get_column_letter(i)].width = min(45, max(8, width + 2))
        ws.freeze_panes = "A5"
        ws.auto_filter.ref = f"A4:{get_column_letter(len(head))}{ws.max_row}"

    ws = wb.active
    ws.title = "Resumen"
    sheet(ws, SUMMARY_HEAD, [_summary(s) for s in summary])
    sheet(wb.create_sheet("Detalle"), DETAIL_HEAD, [_detail(r) for r in rows], status_col=10)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(rows, summary, title, period, company):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1 * cm, rightMargin=1 * cm,
                            topMargin=1 * cm, bottomMargin=1 * cm, title=title)
    st = getSampleStyleSheet()
    small = st["BodyText"].clone("small", fontSize=7, leading=8)

    def table(head, data, widths, status_col=None):
        t = Table([head] + [[Paragraph(str(c), small) for c in line] for line in data],
                  colWidths=widths, repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
                 ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, 0), 7.5),
                 ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
        if status_col is not None:
            for i, line in enumerate(data, 1):
                if line[status_col] in FILLS:
                    style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#" + FILLS[line[status_col]])))
        t.setStyle(TableStyle(style))
        return t

    story = [Paragraph(f"<b>{company or 'Paso Largo'}</b> — {title}", st["Title"]), Paragraph(period, st["Normal"]),
             Spacer(1, 8), Paragraph("Resumen por empleado", st["Heading3"]),
             table(SUMMARY_HEAD, [_summary(s) for s in summary],
                   [1.5 * cm, 5.5 * cm, 3.5 * cm, 2 * cm, 2 * cm, 2 * cm, 1.6 * cm, 2 * cm, 1.8 * cm, 2.4 * cm, 2.2 * cm])]
    if rows:
        story += [PageBreak(), Paragraph("Detalle diario", st["Heading3"]),
                  table(DETAIL_HEAD, [_detail(r) for r in rows],
                        [1.9 * cm, 1 * cm, 1.3 * cm, 4.2 * cm, 3 * cm, 1.4 * cm, 1.4 * cm, 1.5 * cm,
                         1.3 * cm, 1.2 * cm, 2.6 * cm, 6.3 * cm], status_col=10)]
    doc.build(story)
    return buf.getvalue()
