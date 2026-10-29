"""Resumen de locales para el admin: cómo va cada local hoy, en la semana (lunes a hoy) y en el mes (del 1 a hoy).

Todo sale de attendance.compute (una fila por empleado y día) agrupado por el local del empleado (employees.local_id).
Una marca en el biométrico de otro local cuenta para el local del empleado.
"""
import io
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import attendance, db
from .attendance import hm

VISTAS = {"hoy": "Hoy", "semana": "Semana", "mes": "Mes"}
PRESENT = ("Puntual", "Tarde", "Trabajó en día libre")
DOW = "LMMJVSD"
TOP = 5  # cuántos empleados se muestran en «más faltas» y «más tardanzas» del mes


# ------------------------------------------------------------------ fechas
def period(vista: str, today: date) -> tuple[date, date]:
    """Hoy; semana = de lunes a hoy; mes = del 1 a hoy."""
    if vista == "semana":
        return today - timedelta(days=today.weekday()), today
    if vista == "mes":
        return today.replace(day=1), today
    return today, today


def previous_month(start: date, end: date) -> tuple[date, date]:
    """El mes anterior con el mismo número de días (del 1 al 29 → del 1 al 29 del mes anterior, sin pasarse de su fin)."""
    last = start.replace(day=1) - timedelta(days=1)
    first = last.replace(day=1)
    return first, min(first + (end - start), last)


# ------------------------------------------------------------------ datos
def locales() -> list[dict]:
    return db.q("SELECT id, name FROM locales WHERE active=1 ORDER BY name")


def employee_locals() -> dict:
    """{id de empleado: id de su local} (0 si todavía no tiene local)."""
    return {r["id"]: r["local_id"] or 0 for r in db.q("SELECT id, local_id FROM employees")}


def by_local(rows) -> dict:
    """Agrupa las filas de compute por el local del empleado."""
    home, out = employee_locals(), defaultdict(list)
    for r in rows:
        out[home.get(r["employee_id"], 0)].append(r)
    return out


def devices_by_local() -> dict:
    """{local: {'online': n, 'total': n}} de los biométricos activos."""
    out = defaultdict(lambda: {"online": 0, "total": 0})
    for d in db.q("SELECT local_id, online FROM devices WHERE active=1"):
        x = out[d["local_id"] or 0]
        x["total"] += 1
        x["online"] += 1 if d["online"] else 0
    return out


def away_on(day: date) -> dict:
    """Quién marcó ese día en un biométrico de otro local.
    {local: {'out': [nombres de su gente que marcó en otro local], 'in': [nombres de gente de otro local que marcó aquí]}}"""
    out = defaultdict(lambda: {"out": set(), "in": set()})
    for r in db.q("""SELECT DISTINCT e.name, e.local_id AS home, d.local_id AS dev FROM events ev
                     JOIN employees e ON e.employee_no=ev.employee_no JOIN devices d ON d.id=ev.device_id
                     WHERE ev.valid=1 AND e.active=1 AND ev.event_time >= ? AND ev.event_time < ?
                       AND d.local_id IS NOT NULL AND e.local_id IS NOT NULL AND d.local_id<>e.local_id""",
                  (day.isoformat(), (day + timedelta(days=1)).isoformat())):
        out[r["home"]]["out"].add(r["name"])
        out[r["dev"]]["in"].add(r["name"])
    return {k: {"out": sorted(v["out"]), "in": sorted(v["in"])} for k, v in out.items()}


# ------------------------------------------------------------------ cálculos sobre filas de compute
def is_present(r) -> bool:
    return r["status"] in PRESENT or (r["status"] == "Feriado" and bool(r["entry"]))


def counts(rows) -> dict:
    """Totales de un grupo de filas. 'esperados' = quienes debían venir (sin contar permisos ni feriados)."""
    st = Counter(r["status"] for r in rows)
    present = sum(1 for r in rows if is_present(r))
    c = {"presentes": present, "esperados": present + st["Ausente"] + st["Pendiente"],
         "puntual": st["Puntual"], "tarde": st["Tarde"], "ausente": st["Ausente"],
         "justificado": st["Justificado"], "pendiente": st["Pendiente"],
         "worked_min": sum(r["worked_min"] for r in rows), "late_min": sum(r["late_min"] for r in rows),
         "empleados": len({r["employee_id"] for r in rows})}
    c.update(rates(c))
    c["pct_presentes"] = pct(present, c["esperados"])
    return c


def pct(a, b):
    return round(100 * a / b) if b else None


def rates(c) -> dict:
    """% asistencia (vinieron / debían venir, sin contar los de hoy que aún no marcan) y % puntualidad (puntuales / vinieron a horario)."""
    return {"asistencia": pct(c["presentes"], c["presentes"] + c["ausente"]),
            "puntualidad": pct(c["puntual"], c["puntual"] + c["tarde"])}


def delta(now, before):
    """Diferencia en puntos (None si falta un dato)."""
    return None if now is None or before is None else now - before


DOW = "LMMJVSD"


def absences(rows, limit=None) -> list[dict]:
    """Quién faltó, cuántos días y cuáles, de más a menos (con su N° y área para reconocerlo)."""
    days, info = defaultdict(list), {}
    for r in rows:
        if r["status"] == "Ausente":
            days[r["employee_id"]].append(r["date"])
            info[r["employee_id"]] = r
    out = [{"id": eid, "name": info[eid]["name"], "no": info[eid]["employee_no"], "area": info[eid]["department"],
            "dias": len(ds), "fechas": ", ".join(f"{DOW[d.weekday()]} {d.day}" for d in sorted(ds))}
           for eid, ds in days.items()]
    out.sort(key=lambda x: (-x["dias"], x["name"]))
    return out[:limit] if limit else out


def lateness(rows, limit=None) -> list[dict]:
    """Quién llegó tarde más veces (y cuántos minutos en total)."""
    times, mins = Counter(), Counter()
    for r in rows:
        if r["status"] == "Tarde":
            key = (r["employee_id"], r["name"])
            times[key] += 1
            mins[key] += r["late_min"]
    out = [{"name": k[1], "veces": v, "min": mins[k]}
           for k, v in sorted(times.items(), key=lambda x: (-x[1], -mins[x[0]], x[0][1]))]
    return out[:limit] if limit else out


def daily(rows, start: date, end: date) -> list[dict]:
    """Puntuales, tardanzas y faltas de cada día del periodo."""
    by_day = defaultdict(Counter)
    for r in rows:
        by_day[r["date"]][r["status"]] += 1
    out, d = [], start
    while d <= end:
        c = by_day[d]
        out.append({"label": f"{DOW[d.weekday()]} {d.day}", "date": d.isoformat(),
                    "puntual": c["Puntual"] + c["Trabajó en día libre"], "tarde": c["Tarde"], "ausente": c["Ausente"]})
        d += timedelta(days=1)
    return out


CELL = {"Puntual": "ok", "Trabajó en día libre": "ok", "Tarde": "late", "Ausente": "abs", "Justificado": "perm",
        "Feriado": "hol", "Libre": "off", "Pendiente": "pend"}


def week_days(start: date) -> list[date]:
    """Los 7 días de la semana (lunes a domingo) aunque todavía no hayan llegado."""
    return [start + timedelta(days=i) for i in range(7)]


def board(rows, days: list[date], photos: dict) -> list[dict]:
    """Tablero de la semana: una fila por persona y una celda por día (estado + hora de entrada)."""
    people = {}
    for r in rows:
        p = people.setdefault(r["employee_id"], {"id": r["employee_id"], "name": r["name"], "no": r["employee_no"],
                                                 "area": r["department"], "photo": photos.get(r["employee_id"]),
                                                 "by_day": {}, "faltas": 0, "tardes": 0})
        p["by_day"][r["date"]] = r
        p["faltas"] += r["status"] == "Ausente"
        p["tardes"] += r["status"] == "Tarde"
    out = []
    for p in people.values():
        cells = []
        for d in days:
            r = p["by_day"].get(d)
            if not r:
                cells.append({"k": "none", "txt": "", "tip": f"{DOW[d.weekday()]} {d.day}"})
                continue
            entry = r["entry"].strftime("%H:%M") if r["entry"] else ""
            tip = f"{DOW[d.weekday()]} {d.day}: {r['status']}" + (f", entró {entry}" if entry else "") + \
                  (f" ({r['late_min']} min tarde)" if r["late_min"] else "") + (f" · {r['note']}" if r["note"] else "")
            cells.append({"k": CELL.get(r["status"], "ok"), "txt": entry, "late": r["late_min"], "tip": tip})
        out.append({**{k: p[k] for k in ("id", "name", "no", "area", "photo", "faltas", "tardes")}, "cells": cells})
    out.sort(key=lambda x: (-x["faltas"], -x["tardes"], x["name"]))
    return out


def week_strip(rows, days: list[date], today: date) -> list[dict]:
    """Cada día de la semana para toda la empresa: cuántos vinieron a tiempo, tarde o faltaron."""
    by_day = defaultdict(Counter)
    for r in rows:
        by_day[r["date"]][r["status"]] += 1
    out = []
    for d in days:
        c = by_day[d]
        ok, late, absent = c["Puntual"] + c["Trabajó en día libre"], c["Tarde"], c["Ausente"]
        out.append({"dow": "LMMJVSD"[d.weekday()], "day": d.day, "ok": ok, "late": late, "abs": absent,
                    "n": ok + late + absent, "future": d > today, "today": d == today})
    top = max([x["n"] for x in out] + [1])
    for x in out:
        x["h"] = round(100 * x["n"] / top)
    return out


# ------------------------------------------------------------------ armado de cada vista
def build(vista: str = "hoy", now: datetime | None = None) -> dict:
    """Todo lo que muestra /resumen (y exporta a Excel) para la vista pedida."""
    now = now or datetime.now()
    vista = vista if vista in VISTAS else "hoy"
    today = now.date()
    start, end = period(vista, today)
    locs = locales()
    groups = by_local(attendance.compute(start, end, now=now))
    if groups.get(0):
        locs.append({"id": None, "name": "Sin local"})
    keys = [l["id"] or 0 for l in locs]
    shown = [r for k in keys for r in groups.get(k, [])]

    data = {"vista": vista, "title": VISTAS[vista], "start": start, "end": end, "now": now,
            "cards": [], "total": counts(shown)}
    if vista == "hoy":
        devs, away = devices_by_local(), away_on(today)
        for l, k in zip(locs, keys):
            data["cards"].append({**l, **counts(groups.get(k, [])), "devices": devs.get(k, {"online": 0, "total": 0}),
                                  "away": away.get(k, {"out": [], "in": []})})
        data["total"]["devices"] = {"online": sum(d["online"] for d in devs.values()),
                                    "total": sum(d["total"] for d in devs.values())}
        data["total"]["away"] = sum(len(a["out"]) for a in away.values())
    elif vista == "semana":
        days = week_days(start)
        photos = {r["id"]: r["photo"] for r in db.q("SELECT id, photo FROM employees WHERE photo IS NOT NULL")}
        for l, k in zip(locs, keys):
            g = groups.get(k, [])
            data["cards"].append({**l, **counts(g), "faltas": absences(g), "tardanzas": lateness(g),
                                  "series": daily(g, start, end), "board": board(g, days, photos)})
        data["series"] = daily(shown, start, end)
        data["days"] = days
        data["strip"] = week_strip(shown, days, today)
        data["week_no"] = start.isocalendar()[1]
        ranked = sorted((c for c in data["cards"] if c["asistencia"] is not None), key=lambda c: -c["asistencia"])
        data["ranking"] = [{"name": c["name"], "asistencia": c["asistencia"], "puntualidad": c["puntualidad"]} for c in ranked]
    else:
        p1, p2 = previous_month(start, end)
        prev = by_local(attendance.compute(p1, p2, now=now))
        data["prev_start"], data["prev_end"] = p1, p2
        for l, k in zip(locs, keys):
            g, before = counts(groups.get(k, [])), counts(prev.get(k, []))
            data["cards"].append({**l, **g, "prev": before,
                                  "d_asistencia": delta(g["asistencia"], before["asistencia"]),
                                  "d_puntualidad": delta(g["puntualidad"], before["puntualidad"]),
                                  "faltas": absences(groups.get(k, []), TOP), "tardanzas": lateness(groups.get(k, []), TOP)})
        before = counts([r for k in keys for r in prev.get(k, [])])
        data["total"].update(prev=before, d_asistencia=delta(data["total"]["asistencia"], before["asistencia"]),
                             d_puntualidad=delta(data["total"]["puntualidad"], before["puntualidad"]))
    return data


def period_text(data) -> str:
    s, e = data["start"], data["end"]
    if s == e:
        return f"Hoy {s.strftime('%d/%m/%Y')}"
    return f"Del {s.strftime('%d/%m/%Y')} al {e.strftime('%d/%m/%Y')}"


# ------------------------------------------------------------------ Excel
def _p(v):
    return "" if v is None else v / 100


def _sheets(data):
    """[(nombre de la hoja, encabezados, filas, columnas con %)] según la vista."""
    rows = data["cards"] + [{"name": "TOTAL EMPRESA", **data["total"]}]
    if data["vista"] == "hoy":
        head = ["Local", "Debían venir", "Se presentaron", "% presentes", "Puntuales", "Tarde", "Ausentes",
                "Con permiso", "Pendientes", "Biométricos en línea", "Biométricos", "Marcaron en otro local"]
        body = [[c["name"], c["esperados"], c["presentes"], _p(c["pct_presentes"]), c["puntual"], c["tarde"],
                 c["ausente"], c["justificado"], c["pendiente"], c["devices"]["online"], c["devices"]["total"],
                 len(c["away"]["out"]) if isinstance(c.get("away"), dict) else c.get("away", 0)] for c in rows]
        return [("Hoy", head, body, {4})]
    if data["vista"] == "semana":
        head = ["Local", "Asistencias", "Puntuales", "Tarde", "Faltas", "Con permiso", "% asistencia", "% puntualidad",
                "Horas trabajadas"]
        body = [[c["name"], c["presentes"], c["puntual"], c["tarde"], c["ausente"], c["justificado"],
                 _p(c["asistencia"]), _p(c["puntualidad"]), hm(c["worked_min"])] for c in rows]
        days = [["Local", "Día", "Puntual", "Tarde", "Ausente"]]
        days_body = [[c["name"], s["date"], s["puntual"], s["tarde"], s["ausente"]] for c in data["cards"] for s in c["series"]]
        return [("Semana", head, body, {7, 8}), ("Por día", days[0], days_body, set()),
                ("Faltas", ["Local", "N°", "Empleado", "Área", "Días que faltó", "Cuáles"],
                 [[c["name"], f["no"], f["name"], f["area"], f["dias"], f["fechas"]] for c in data["cards"]
                  for f in c["faltas"]], set()),
                ("Tardanzas", ["Local", "Empleado", "Veces tarde", "Minutos tarde"],
                 [[c["name"], t["name"], t["veces"], t["min"]] for c in data["cards"] for t in c["tardanzas"]], set())]
    head = ["Local", "% asistencia", "% asistencia mes anterior", "% puntualidad", "% puntualidad mes anterior",
            "Faltas", "Tardanzas", "Con permiso", "Horas trabajadas"]
    body = [[c["name"], _p(c["asistencia"]), _p(c["prev"]["asistencia"]), _p(c["puntualidad"]),
             _p(c["prev"]["puntualidad"]), c["ausente"], c["tarde"], c["justificado"], hm(c["worked_min"])] for c in rows]
    return [("Mes", head, body, {2, 3, 4, 5}),
            (f"Top {TOP} faltas", ["Local", "Empleado", "Días que faltó"],
             [[c["name"], f["name"], f["dias"]] for c in data["cards"] for f in c["faltas"]], set()),
            (f"Top {TOP} tardanzas", ["Local", "Empleado", "Veces tarde", "Minutos tarde"],
             [[c["name"], t["name"], t["veces"], t["min"]] for c in data["cards"] for t in c["tardanzas"]], set())]


def to_excel(data, company="") -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4A5726")
    title = f"Resumen de locales · {data['title']}" + (f" · {company}" if company else "")
    sub = period_text(data)
    if data["vista"] == "mes":
        sub += f" (comparado con {data['prev_start'].strftime('%d/%m')} al {data['prev_end'].strftime('%d/%m/%Y')})"
    for name, head, body, pct_cols in _sheets(data):
        ws = wb.create_sheet(name[:31])
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([sub])
        ws.append([])
        ws.append(head)
        for i in range(1, len(head) + 1):
            cell = ws.cell(row=4, column=i)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in body:
            ws.append(row)
            if row and row[0] == "TOTAL EMPRESA":
                for cell in ws[ws.max_row]:
                    cell.font = Font(bold=True)
            for col in pct_cols:
                ws.cell(row=ws.max_row, column=col).number_format = "0%"
        if not body:
            ws.append(["Sin datos en este periodo"])
        for i, h in enumerate(head, 1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in body if i - 1 < len(r)])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 40)
        ws.freeze_panes = "B5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
